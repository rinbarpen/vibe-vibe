"""Caption adjacency, explicit keep chains and bounded native-layout evidence."""
from copy import deepcopy
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from PIL import Image
except ImportError as exc:
    raise unittest.SkipTest("Optional python-docx/Pillow are not installed") from exc

REPO = Path(__file__).resolve().parents[3]
SCRIPTS = REPO / "manifests/office/docx/scripts"


def load(name):
    spec = importlib.util.spec_from_file_location("office_caption_test_" + name, SCRIPTS / (name + ".py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


eq = load("equations")
native = load("native_format")
W, M = eq.W, eq.M
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"


def flag(paragraph, name):
    item = paragraph.find(W + "pPr/" + W + name)
    return None if item is None else item.get(W + "val", "true").lower() not in {"0", "false", "off"}


def uno_available():
    if not (shutil.which("libreoffice") or shutil.which("soffice")):
        return False
    try:
        return Path("/usr/bin/python3").is_file() and subprocess.run(
            ["/usr/bin/python3", "-c", "import uno"], capture_output=True, timeout=5).returncode == 0
    except (OSError, subprocess.TimeoutExpired):
        return False


class CaptionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="office caption tests ")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.picture = BytesIO()
        Image.new("RGB", (120, 50), "teal").save(self.picture, format="PNG")
        self.picture.seek(0)

    def document(self):
        document = Document()
        document.add_picture(self.picture, width=Inches(3))
        figure = eq.add_caption(document, "Figure", "图下标题", bookmark="FigureExample")
        caption = eq.add_caption(document, "Table", "表上标题", bookmark="TableExample")
        table = document.add_table(rows=3, cols=2)
        for index, row in enumerate(table.rows):
            for column, cell in enumerate(row.cells):
                cell.text = f"值 {index}-{column}"
        return document, figure["paragraph"], caption["paragraph"], table

    def test_real_seq_image_below_and_caption_above_whole_table_have_terminated_chains(self):
        document, figure, caption, table = self.document()
        following = document.add_paragraph("后续正文不属于表格组")
        before_order = list(document.element.body)
        report = eq.apply_caption_layout(document)
        self.assertEqual(report["status"], "caption_layout_applied", report)
        self.assertEqual(report["quality_status"], "pending")
        self.assertFalse(report["pages_verified"])
        self.assertEqual([item["sequence"] for item in report["bindings"]], ["Figure", "Table"])
        picture = figure._p.getprevious()
        self.assertTrue(flag(picture, "keepNext"))
        self.assertTrue(flag(picture, "keepLines"))
        self.assertTrue(flag(figure._p, "keepLines"))
        self.assertFalse(flag(figure._p, "keepNext"))
        self.assertTrue(flag(caption._p, "keepLines"))
        self.assertTrue(flag(caption._p, "keepNext"))
        for index, row in enumerate(table.rows):
            self.assertIsNotNone(row._tr.find(W + "trPr/" + W + "cantSplit"))
            for paragraph in row._tr.iter(W + "p"):
                self.assertEqual(flag(paragraph, "keepNext"), index < len(table.rows) - 1)
        self.assertIsNone(flag(following._p, "keepNext"))
        self.assertEqual(list(document.element.body), before_order)
        self.assertEqual(eq.validate_references(document)["status"], "references_valid")

    def test_add_caption_immediately_binds_existing_picture_and_table_waits_for_object(self):
        document = Document()
        document.add_picture(self.picture)
        figure = eq.add_caption(document, "Figure", "图在上题注在下")
        self.assertEqual(figure["layout"]["bindings"][0]["sequence"], "Figure")
        caption = eq.add_caption(document, "Table", "随后再建表")
        self.assertEqual(caption["layout"]["status"], "pending")
        table = document.add_table(rows=2, cols=2)
        report = eq.apply_caption_layout(document)
        self.assertEqual(report["status"], "caption_layout_applied", report)
        self.assertEqual(len(report["bindings"]), 2)
        self.assertFalse(flag(next(table.rows[-1]._tr.iter(W + "p")), "keepNext"))

    def test_explicit_multipage_policy_clears_cross_row_chain_only_when_requested(self):
        document, _, caption, table = self.document()
        eq.apply_caption_layout(document)
        report = eq.apply_caption_layout(document, allow_multipage_tables=True)
        self.assertTrue(report["allow_multipage_tables"])
        self.assertFalse(report["bindings"][1]["whole_table_kept"])
        self.assertTrue(flag(caption._p, "keepNext"))
        self.assertTrue(all(flag(p, "keepNext") is False for p in table._tbl.iter(W + "p")))
        before = document.element.xml
        with self.assertRaises(ValueError):
            eq.apply_caption_layout(document, allow_multipage_tables="yes")
        self.assertEqual(document.element.xml, before)

    def test_plain_text_labels_are_not_caption_semantics(self):
        document = Document()
        document.add_paragraph("图 1 只是正文")
        document.add_paragraph("表 1 不是SEQ题注")
        document.add_table(rows=2, cols=2)
        before = document.element.xml
        report = eq.apply_caption_layout(document)
        self.assertEqual(report["bindings"], [])
        self.assertEqual(report["pending"], [])
        self.assertEqual(document.element.xml, before)

    def test_nonadjacent_floating_and_explicit_boundaries_remain_pending_without_moves(self):
        for kind in ("empty_gap", "floating", "page_break", "section_break"):
            document = Document()
            document.add_picture(self.picture)
            if kind == "empty_gap":
                document.add_paragraph()
            caption = eq.add_caption(document, "Figure", "需检查", bookmark="FigurePending")["paragraph"]
            picture = document.paragraphs[0]._p
            if kind == "floating":
                next(picture.iter(WP + "inline")).tag = WP + "anchor"
            elif kind == "page_break":
                caption.paragraph_format.page_break_before = True
            elif kind == "section_break":
                picture.get_or_add_pPr().append(OxmlElement("w:sectPr"))
            before = document.element.xml
            report = eq.apply_caption_layout(document)
            with self.subTest(kind=kind):
                self.assertEqual(report["status"], "pending", report)
                self.assertTrue(report["pending"])
                self.assertEqual(document.element.xml, before)
            self.picture.seek(0)

    def test_sdt_caption_and_equation_layout_are_not_repositioned_as_business_tables(self):
        document = Document()
        caption = eq.add_caption(document, "Table", "不能绑到公式骨架", bookmark="TablePending")["paragraph"]
        table = document.add_table(rows=1, cols=3)
        marker = OxmlElement("w:tblCaption")
        marker.set(W + "val", "OfficeEquationLayout")
        table._tbl.tblPr.append(marker)
        before = document.element.xml
        report = eq.apply_caption_layout(document)
        self.assertEqual(report["pending"][0]["reason"], "equation_layout_is_not_a_business_table")
        self.assertEqual(document.element.xml, before)
        wrapper, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        caption._p.addprevious(wrapper)
        wrapper.append(content)
        content.append(caption._p)
        before = document.element.xml
        report = eq.apply_caption_layout(document)
        self.assertEqual(report["pending"][0]["reason"], "caption_outside_direct_main_body")
        self.assertEqual(document.element.xml, before)

    def test_last_row_all_paragraphs_end_chain_even_when_template_inherits_keepnext(self):
        document, _, _, table = self.document()
        document.styles["Normal"].paragraph_format.keep_with_next = True
        table.rows[-1].cells[0].add_paragraph("最后一行第二段")
        eq.apply_caption_layout(document)
        for p in table.rows[-1]._tr.iter(W + "p"):
            self.assertFalse(flag(p, "keepNext"))
        first = next(table.rows[0]._tr.iter(W + "p"))
        self.assertTrue(flag(first, "keepNext"))

    def test_oversized_minimum_geometry_is_pending_without_crop_or_font_changes(self):
        document, figure, _, table = self.document()
        for row in table.rows:
            row.height, row.height_rule = Inches(4), WD_ROW_HEIGHT_RULE.AT_LEAST
        height = next(figure._p.getprevious().iter(WP + "extent"))
        height.set("cy", str(Inches(20)))
        before_fonts = [node.xml for node in document.element.iter(W + "rPr")]
        before_extent = deepcopy(height.attrib)
        report = eq.apply_caption_layout(document)
        self.assertEqual(report["status"], "pending")
        reasons = {item["reason"] for item in report["pending"]}
        self.assertIn("figure_exceeds_declared_page_body", reasons)
        self.assertIn("minimum_table_height_exceeds_declared_page_body", reasons)
        self.assertEqual(height.attrib, before_extent)
        self.assertEqual([node.xml for node in document.element.iter(W + "rPr")], before_fonts)
        self.assertFalse(report["pages_verified"])

    @unittest.skipUnless(shutil.which("pandoc"), "Pandoc is not installed")
    def test_formula_blocks_bookmarks_and_ref_instructions_are_preserved(self):
        document, _, _, _ = self.document()
        equation = eq.add_numbered_equation(document, "E=mc^2", bookmark="Energy")
        eq.add_ref_field(document.add_paragraph("参见公式 "), "Energy")
        before = equation["table"]._tbl.xml
        fields = [node.get(W + "instr") for node in document.element.iter(W + "fldSimple")]
        eq.apply_caption_layout(document)
        self.assertEqual(equation["table"]._tbl.xml, before)
        self.assertEqual([node.get(W + "instr") for node in document.element.iter(W + "fldSimple")], fields)
        self.assertEqual(eq.validate_references(document)["status"], "references_valid")

    def test_native_profile_and_word_macro_have_explicit_default_and_chain_termination(self):
        self.assertIs(native.DEFAULT_PROFILE["allow_multipage_tables"], False)
        profile = self.root / "profile.json"
        profile.write_text('{"allow_multipage_tables":true}')
        self.assertTrue(native.profile_settings(profile)["allow_multipage_tables"])
        macro = (REPO / "manifests/office/docx/macros/OfficeLayout.bas").read_text()
        for token in ("Public Sub ApplyOfficeCaptionLayout()", "ALLOW_MULTIPAGE_TABLES As Boolean = False",
                      "item.Type = wdFieldSequence", 'candidate = "FIGURE" Or candidate = "TABLE"',
                      "previous.Range.ParagraphFormat.KeepWithNext = True",
                      "caption.Range.ParagraphFormat.KeepWithNext = False",
                      "target.Range.ParagraphFormat.KeepWithNext = Not allowMultipage",
                      "finalRow.Range.ParagraphFormat.KeepWithNext = False",
                      "target.Rows.AllowBreakAcrossPages = False", "GroupSpansPages(combined)"):
            self.assertIn(token, macro)
        executable = "\n".join(line for line in macro.splitlines() if not line.lstrip().startswith("'"))
        self.assertNotIn("SaveAs", executable)
        self.assertNotIn("Sub AutoOpen", executable)
        self.assertNotIn("Sub Document_Open", executable)
        source = (SCRIPTS / "native_format.py").read_text()
        self.assertIn("neighbour.Split = allow_multipage_tables", source)
        self.assertIn("neighbour.KeepTogether = False", source)
        self.assertIn('caption.ParaKeepTogether, caption.ParaSplit = False, False', source)

    @unittest.skipUnless(uno_available(), "LibreOffice with system Python UNO is not installed")
    def test_real_uno_binds_native_seq_and_preserves_explicit_end_flags(self):
        document, _, _, _ = self.document()
        source, output = self.root / "before.docx", self.root / "after.docx"
        document.save(source)
        before = native.sha256(source)
        report = native.format_document(source, output)
        self.assertEqual(report["status"], "formatted", report)
        layout = report["native_changes"]["caption_layout"]
        self.assertEqual(layout["status"], "caption_layout_applied", layout)
        self.assertEqual(len(layout["bindings"]), 2)
        self.assertFalse(layout["pages_verified"])
        self.assertEqual(native.sha256(source), before)
        result = Document(output)
        paragraphs = result.paragraphs
        self.assertTrue(paragraphs[0].paragraph_format.keep_with_next)
        self.assertFalse(paragraphs[1].paragraph_format.keep_with_next)
        self.assertTrue(paragraphs[2].paragraph_format.keep_with_next)
        rows = result.tables[0].rows
        self.assertTrue(rows[0].cells[0].paragraphs[0].paragraph_format.keep_with_next)
        self.assertFalse(rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next)

    @unittest.skipUnless(uno_available(), "LibreOffice with system Python UNO is not installed")
    def test_real_uno_oversized_group_withholds_output_and_keeps_readable_source(self):
        document = Document()
        eq.add_caption(document, "Table", "超高表", bookmark="TooTall")
        table = document.add_table(rows=8, cols=2)
        for row in table.rows:
            row.height, row.height_rule = Inches(2), WD_ROW_HEIGHT_RULE.AT_LEAST
        source, output = self.root / "oversize.docx", self.root / "withheld.docx"
        document.save(source)
        report = native.format_document(source, output)
        self.assertEqual(report["status"], "failed", report)
        self.assertTrue(report["input_unchanged"])
        self.assertFalse(output.exists())
        layout = report["native_changes"]["caption_layout"]
        self.assertTrue(layout["blocking_pending"])
        self.assertEqual(layout["pending"][0]["reason"], "minimum_table_height_exceeds_declared_page_body")
        self.assertEqual(report["quality_status"], "pending")


if __name__ == "__main__":
    unittest.main()
