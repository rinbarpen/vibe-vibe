"""Native Word-layout helper guards, preservation gates and optional real UNO checks."""
import importlib.util
import json
import os
from pathlib import Path
import shutil
import signal
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch
import xml.etree.ElementTree as ET
import zipfile


REPO = Path(__file__).resolve().parents[3]
SCRIPT = REPO / "manifests/office/docx/scripts/native_format.py"
MACRO = REPO / "manifests/office/docx/macros/OfficeLayout.bas"
spec = importlib.util.spec_from_file_location("office_native_format_test_module", SCRIPT)
native = importlib.util.module_from_spec(spec)
spec.loader.exec_module(native)
W = native.W
NS = {"w": W[1:-1]}


def engine_available():
    if not (shutil.which("libreoffice") or shutil.which("soffice")):
        return False
    for candidate in dict.fromkeys((sys.executable, "/usr/bin/python3", shutil.which("python3"))):
        if candidate and Path(candidate).is_file():
            try:
                if subprocess.run([candidate, "-c", "import uno"], capture_output=True, timeout=5).returncode == 0:
                    return True
            except (OSError, subprocess.TimeoutExpired):
                pass
    return False


class NativeFormatTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="office native format ")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def package(self, name="source.docx", content=None):
        path = self.root / name
        body = content or '<w:p><w:r><w:t>原文未改</w:t></w:r></w:p>'
        with zipfile.ZipFile(path, "w") as package:
            package.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
            package.writestr("word/document.xml", f'<w:document xmlns:w="{W[1:-1]}"><w:body>{body}</w:body></w:document>')
        return path

    def cli(self, *args):
        return subprocess.run([sys.executable, str(SCRIPT), "--engine", "libreoffice", *map(str, args)],
                              capture_output=True, text=True, timeout=90,
                              env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})

    def test_explicit_working_copy_and_existing_output_guards(self):
        source = self.package()
        before = native.sha256(source)
        with self.assertRaisesRegex(ValueError, "different working-copy"):
            native.format_document(source, source)
        output = self.root / "existing.docx"
        output.write_bytes(b"preserve output")
        with self.assertRaisesRegex(ValueError, "Output exists"):
            native.format_document(source, output)
        self.assertEqual(output.read_bytes(), b"preserve output")
        self.assertEqual(native.sha256(source), before)

    def test_docm_and_embedded_macro_packages_are_not_silently_converted(self):
        source = self.package("source.docm")
        with self.assertRaisesRegex(ValueError, "DOCM/DOTM"):
            native.format_document(source, self.root / "output.docx")
        source = self.package()
        with zipfile.ZipFile(source, "a") as package:
            package.writestr("word/vbaProject.bin", b"macro sentinel")
        before = source.read_bytes()
        result = native.format_document(source, self.root / "output.docx")
        self.assertEqual(result["status"], "failed")
        self.assertIn("Macro-bearing", result["error"])
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse((self.root / "output.docx").exists())

    def test_external_data_fields_are_not_refreshed_by_native_worker(self):
        for command in ("DDE", "DDEAUTO", "DATABASE", "LINK", "INCLUDETEXT", "INCLUDEPICTURE", "IMPORT", "RD", "ADDIN", "CITATION", "BIBLIOGRAPHY"):
            with self.subTest(command=command):
                source = self.package(content=f'<w:p><w:fldSimple w:instr="{command} external-source"/></w:p>')
                before = native.sha256(source)
                with patch.object(native.subprocess, "Popen") as launch:
                    result = native.format_document(source, self.root / "output.docx")
                launch.assert_not_called()
                self.assertEqual(result["status"], "failed")
                self.assertIn("selective Word updates", result["error"])
                self.assertEqual(result["external_fields_not_updated"], [command + " external-source"])
                self.assertTrue(result["input_unchanged"])
                self.assertEqual(native.sha256(source), before)
                self.assertFalse((self.root / "output.docx").exists())

    def test_profile_validation_and_explicit_overrides(self):
        profile = self.root / "profile.json"
        profile.write_text(json.dumps({"font": "Example Family", "body_size_pt": 12}))
        result = native.profile_settings(profile)
        self.assertEqual(result["font"], "Example Family")
        self.assertEqual(result["body_size_pt"], 12)
        self.assertEqual(result["heading_sizes_pt"], [18, 15, 13])
        for value in ({"trust_macros": True}, {"font": ""}, {"body_size_pt": True},
                      {"heading_sizes_pt": [10, 12]}, {"short_row_max_chars": 0},
                      {"repeat_table_headers": 1}, {"caption_size_pt": float("nan")}):
            with self.subTest(value=value):
                profile.write_text(json.dumps(value))
                with self.assertRaises(ValueError):
                    native.profile_settings(profile)

    def test_report_source_profile_and_output_collisions_are_rejected(self):
        source = self.package()
        before = source.read_bytes()
        profile = self.root / "profile.json"
        profile.write_text("{}")
        output = self.root / "output.docx"
        for report in (source, output, profile):
            with self.subTest(report=report):
                result = self.cli("--input", source, "--output", output, "--report", report, "--profile", profile)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("Report path must", result.stderr)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(profile.read_text(), "{}")
        self.assertFalse(output.exists())

    def test_validation_failure_still_writes_requested_new_report(self):
        source = self.package()
        report = self.root / "validation.json"
        result = self.cli("--input", source, "--output", self.root / "output.docx", "--report", report, "--timeout", 0)
        self.assertEqual(result.returncode, 1)
        data = json.loads(report.read_text())
        self.assertEqual(data["status"], "failed")
        self.assertTrue(data["input_unchanged"])
        self.assertIn("positive finite", data["error"])

    def test_missing_engine_reports_pending_and_source_unchanged(self):
        source = self.package()
        with patch.object(native.shutil, "which", return_value=None):
            result = native.format_document(source, self.root / "output.docx")
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["quality_status"], "pending")
        self.assertIn("Missing LibreOffice", result["error"])
        self.assertTrue(result["input_unchanged"])
        self.assertEqual(result["commands"], [])

    def test_timeout_records_command_and_kills_only_owned_group(self):
        source = self.package()
        process = Mock(pid=424242, returncode=-signal.SIGTERM)
        process.poll.return_value = None
        process.wait.side_effect = [subprocess.TimeoutExpired("owned-office", 3), 0]
        run_values = [subprocess.CompletedProcess([], 0, "Fixture Office", ""),
                      subprocess.TimeoutExpired("worker", 1, output=b"partial", stderr=b"stalled")]
        with patch.object(native.shutil, "which", return_value="/fixture/libreoffice"), \
             patch.object(native, "uno_python", return_value="/fixture/python"), \
             patch.object(native.subprocess, "run", side_effect=run_values), \
             patch.object(native.subprocess, "Popen", return_value=process) as launch, \
             patch.object(native.os, "killpg") as kill:
            report = native.format_document(source, self.root / "output.docx", timeout=1)
        self.assertEqual(report["status"], "failed")
        self.assertTrue(report["input_unchanged"])
        self.assertFalse((self.root / "output.docx").exists())
        self.assertTrue(report["commands"][-1]["timed_out"])
        self.assertEqual(report["commands"][-1]["stdout"], "partial")
        kill.assert_called_once_with(424242, signal.SIGTERM)
        self.assertTrue(launch.call_args.kwargs["start_new_session"])
        command = launch.call_args.args[0]
        self.assertTrue(any(arg.startswith("--accept=pipe,name=office_native_") for arg in command))
        self.assertTrue(any(arg.startswith("-env:UserInstallation=file:") for arg in command))

    def test_preservation_inventory_catches_fields_locks_native_objects_and_text(self):
        xml = '''<w:p><w:bookmarkStart w:id="1" w:name="FigureOne"/><w:r><w:t>图 </w:t></w:r>
          <w:fldSimple w:instr=" SEQ Figure "><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>
          <w:sdt><w:sdtPr><w:tag w:val="office.title"/><w:lock w:val="sdtContentLocked"/>
          <w:text/></w:sdtPr><w:sdtContent><w:p><w:r><w:t>标题</w:t></w:r></w:p></w:sdtContent></w:sdt>
          <w:p><w:r><w:drawing/></w:r></w:p><w:tbl/>
          <w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> REF FigureOne \\h </w:instrText></w:r>
          <w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>1</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'''
        source = self.package(content=xml)
        baseline = native.inspect_package(source)
        self.assertEqual(baseline["fields"], ["SEQ Figure", "REF FigureOne \\h"])
        self.assertIn("sdtContentLocked", baseline["controls"][0])
        self.assertEqual(baseline["text"], ["图", "标题"])
        self.assertIn(W + "tbl", baseline["objects"])
        for key in ("fields", "controls", "bookmarks", "objects", "text"):
            with self.subTest(key=key):
                after = {k: list(v) for k, v in baseline.items()}
                after[key] = []
                self.assertTrue(any(key in error or key == "text" and "text" in error
                                    for error in native.preservation_errors(baseline, after)))
        self.assertEqual(native.preservation_errors(baseline, baseline), [])

    def test_atomic_report_publication_does_not_replace_existing_evidence(self):
        report = self.root / "report.json"
        native.write_report(report, {"status": "first"})
        with self.assertRaises(FileExistsError):
            native.write_report(report, {"status": "second"})
        self.assertEqual(json.loads(report.read_text())["status"], "first")

    def test_macro_is_explicit_native_vba_without_trust_changes_or_autorun(self):
        macro = MACRO.read_text()
        for entry in ("ApplyOfficeLayout", "UpdateOfficeFields", "ProtectSelectedTemplateControl", "RemoveSelectedTemplateControlLocks"):
            self.assertIn("Public Sub " + entry + "()", macro)
        for token in ("wdStyleNormal", "wdStyleHeading1", ".HeadingFormat = True", ".KeepWithNext = True",
                      ".AllowBreakAcrossPages = False", "doc.Repaginate", "doc.StoryRanges", "item.Locked",
                      "control.LockContents = (choice = vbYes)", "control.LockContentControl = True"):
            self.assertIn(token, macro)
        executable = "\n".join(line for line in macro.splitlines() if not line.lstrip().startswith("'"))
        for token in ("Sub AutoOpen", "Sub Document_Open", "VBProject", "NormalTemplate", "SaveAs", ".Save", "AcceptAllRevisions", "TrackRevisions =", "AutomationSecurity", "Shell("):
            self.assertNotIn(token, executable)

    def test_uno_worker_uses_symbolic_macro_never_execute_on_both_loads(self):
        source = SCRIPT.read_text()
        self.assertIn('uno.getConstantByName("com.sun.star.document.MacroExecMode.NEVER_EXECUTE")', source)
        self.assertEqual(source.count("MacroExecutionMode=never_execute"), 2)
        self.assertNotRegex(source, r"MacroExecutionMode\s*=\s*[1-9]")

    @unittest.skipUnless(engine_available(), "LibreOffice with Python UNO is not installed")
    def test_real_native_engine_formats_reopens_and_preserves_word_features(self):
        try:
            from docx import Document
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError:
            self.skipTest("python-docx is not installed")
        source = self.root / "原生 输入.docx"
        output = self.root / "原生 输出.docx"
        doc = Document()
        doc.add_heading("中文标题", 1)
        doc.add_paragraph("正文保持不变，并保留原生表格和域。")
        doc.add_paragraph("图 1 原生题注", style="Caption")
        table = doc.add_table(rows=3, cols=2)
        for index, row in enumerate(table.rows):
            for column, cell in enumerate(row.cells):
                cell.text = f"第 {index} 行 {column} 列"
        # Long cells should remain splittable, rather than blocking a whole page.
        table.rows[2].cells[0].text = "长行内容" * 100
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), " PAGE ")
        doc.add_paragraph("页码：")._p.append(field)
        bookmark = OxmlElement("w:bookmarkStart")
        bookmark.set(qn("w:id"), "1")
        bookmark.set(qn("w:name"), "OfficeBody")
        doc.paragraphs[1]._p.insert(0, bookmark)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "1")
        doc.paragraphs[1]._p.append(end)
        merged = doc.add_table(rows=3, cols=3)
        for index, row in enumerate(merged.rows):
            for column, cell in enumerate(row.cells):
                cell.text = f"合并单元格 {index},{column}"
        merged.cell(0, 0).merge(merged.cell(0, 1))
        merged.cell(1, 0).merge(merged.cell(2, 0))
        doc.save(source)
        before = native.sha256(source)
        report = native.format_document(source, output, timeout=60)
        self.assertEqual(report["status"], "formatted", report)
        self.assertTrue(report["input_unchanged"])
        self.assertTrue(report["preservation_passed"])
        self.assertTrue(report["native_changes"]["native_reopen_passed"])
        self.assertEqual(report["native_changes"]["field_refresh_passes"], 2)
        self.assertEqual(report["native_changes"]["tables"], 2)
        self.assertEqual(native.sha256(source), before)
        self.assertEqual(report["quality_status"], "pending")
        with zipfile.ZipFile(output) as package:
            styles = ET.fromstring(package.read("word/styles.xml"))
            main = ET.fromstring(package.read("word/document.xml"))
        heading = next(s for s in styles.findall("w:style", NS) if s.get(W + "styleId") == "Heading1")
        self.assertIsNotNone(heading.find("w:pPr/w:keepNext", NS))
        self.assertEqual(heading.find("w:rPr/w:sz", NS).get(W + "val"), "36")
        rows = main.findall(".//w:tr", NS)
        self.assertIsNotNone(rows[0].find("w:trPr/w:tblHeader", NS))
        self.assertIsNotNone(rows[0].find("w:trPr/w:cantSplit", NS))
        self.assertIsNone(rows[2].find("w:trPr/w:cantSplit", NS))
        self.assertIn("PAGE", native.inspect_package(output)["fields"])
        self.assertIn("OfficeBody", native.inspect_package(output)["bookmarks"])


if __name__ == "__main__":
    unittest.main()
