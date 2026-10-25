"""Real LaTeX fragments, native equation/number fields and preservation gates."""
from copy import deepcopy
import hashlib
import importlib.util
from io import BytesIO
from lxml import etree
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError as exc:
    raise unittest.SkipTest("Optional python-docx is not installed") from exc

REPO = Path(__file__).resolve().parents[3]
SCRIPTS = REPO / "manifests/office/docx/scripts"


def load(name):
    spec = importlib.util.spec_from_file_location("office_equation_test_" + name, SCRIPTS / (name + ".py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


eq = load("equations")
native = load("native_format")
M, W = eq.M, eq.W


class EquationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="office equations ")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_independent_sequences_dirty_caches_and_paired_bookmarks(self):
        document = Document()
        for sequence in ("Figure", "Table", "Equation"):
            for suffix in ("A", "B"):
                name = sequence + suffix
                self.assertEqual(eq.add_seq_field(document.add_paragraph(), sequence, name), name)
        actual = eq.validate_references(document)
        self.assertEqual(actual["status"], "references_valid", actual)
        self.assertEqual(actual["sequences"], {"Figure": 2, "Table": 2, "Equation": 2})
        self.assertEqual(actual["quality_status"], "pending")
        starts = list(document.element.iter(W + "bookmarkStart"))
        ends = list(document.element.iter(W + "bookmarkEnd"))
        self.assertEqual(len({item.get(W + "id") for item in starts}), 6)
        self.assertEqual({item.get(W + "id") for item in starts}, {item.get(W + "id") for item in ends})
        for item in document.element.iter(W + "fldSimple"):
            self.assertEqual(item.get(W + "dirty"), "true")
            self.assertEqual("".join(n.text or "" for n in item.iter(W + "t")), "?")
        self.assertEqual(document.settings.element.find(W + "updateFields").get(W + "val"), "true")

    def test_bookmark_wraps_only_number_field_not_caption_text(self):
        document = Document()
        result = eq.add_caption(document, "Figure", "月度销售", bookmark="FigureSales")
        nodes = list(result["paragraph"]._p)
        start = next(i for i, node in enumerate(nodes) if node.tag == W + "bookmarkStart")
        end = next(i for i, node in enumerate(nodes) if node.tag == W + "bookmarkEnd")
        self.assertEqual([node.tag for node in nodes[start + 1:end]], [W + "fldSimple"])
        self.assertEqual(nodes[start + 1].get(W + "instr").strip(), "SEQ Figure \\* ARABIC")
        self.assertIn("图 ", "".join(nodes[start - 1].itertext()))
        self.assertIn("月度销售", "".join(nodes[end + 1].itertext()))

    def test_invalid_names_sequences_and_duplicate_casing_leave_paragraph_unchanged(self):
        document = Document()
        paragraph = document.add_paragraph("保留内容")
        for sequence, name in (("FigureX", "ValidName"), ("Equation", "3bad"),
                               ("Table", "has space"), ("Table", "has-hyphen"),
                               ("Table", "x" * 41), ("Table", "x\\h")):
            before = paragraph._p.xml
            with self.subTest(sequence=sequence, name=name), self.assertRaises(ValueError):
                eq.add_seq_field(paragraph, sequence, name)
            self.assertEqual(paragraph._p.xml, before)
        eq.add_seq_field(paragraph, "Table", "StableTable")
        before = paragraph._p.xml
        with self.assertRaisesRegex(ValueError, "Duplicate bookmark"):
            eq.add_seq_field(paragraph, "Table", "stabletable")
        self.assertEqual(paragraph._p.xml, before)
        before = document.element.xml
        with self.assertRaises(ValueError):
            eq.add_caption(document, "Figure", "重复", bookmark="StableTable")
        self.assertEqual(document.element.xml, before)

    def test_generated_names_do_not_encode_ordinal_and_survive_reordering(self):
        document = Document()
        one, two = document.add_paragraph(), document.add_paragraph()
        first = eq.add_seq_field(one, "Figure")
        second = eq.add_seq_field(two, "Figure")
        self.assertNotEqual(first, second)
        self.assertRegex(first, r"^OfficeFigure_[0-9a-f]{16}$")
        one._p.addprevious(two._p)
        report = eq.validate_references(document)
        self.assertEqual(set(report["bookmark_names"]), {first, second})
        self.assertEqual(report["status"], "references_valid")

    def test_missing_and_forward_ref_targets_are_explicit_and_finally_validated(self):
        document = Document()
        reference = document.add_paragraph("如图 ")
        before = reference._p.xml
        with self.assertRaisesRegex(ValueError, "does not exist"):
            eq.add_ref_field(reference, "FutureFigure")
        self.assertEqual(reference._p.xml, before)
        field = eq.add_ref_field(reference, "FutureFigure", allow_missing=True)
        self.assertEqual(field.get(W + "instr").strip(), "REF FutureFigure \\h")
        self.assertEqual(eq.validate_references(document)["status"], "failed")
        eq.add_seq_field(document.add_paragraph(), "Figure", "FutureFigure")
        report = eq.validate_references(document)
        self.assertEqual(report["status"], "references_valid", report)
        self.assertEqual(report["quality_status"], "pending")
        self.assertIn("unverified", report["field_results"])

    def test_existing_numeric_id_or_pair_corruption_is_not_overwritten(self):
        for kind in ("bad-id", "duplicate-id", "unpaired"):
            document = Document()
            paragraph = document.add_paragraph()
            eq.add_seq_field(paragraph, "Figure", "FigureA")
            start = next(document.element.iter(W + "bookmarkStart"))
            if kind == "bad-id":
                start.set(W + "id", "not-an-id")
            elif kind == "duplicate-id":
                duplicate = deepcopy(start)
                duplicate.set(W + "name", "OtherName")
                paragraph._p.append(duplicate)
            else:
                end = next(document.element.iter(W + "bookmarkEnd"))
                end.getparent().remove(end)
            before = paragraph._p.xml
            with self.subTest(kind=kind), self.assertRaises(ValueError):
                eq.add_seq_field(paragraph, "Figure", "FigureB")
            self.assertEqual(paragraph._p.xml, before)
            self.assertEqual(eq.validate_references(document)["status"], "failed")

    def test_explicit_cache_is_still_dirty_and_unverified(self):
        document = Document()
        paragraph = document.add_paragraph()
        eq.add_seq_field(paragraph, "Equation", "Energy", cached_text="42")
        field = next(paragraph._p.iter(W + "fldSimple"))
        self.assertEqual("".join(n.text or "" for n in field.iter(W + "t")), "42")
        self.assertEqual(field.get(W + "dirty"), "true")
        self.assertEqual(eq.validate_references(document)["quality_status"], "pending")
        before = paragraph._p.xml
        with self.assertRaises(ValueError):
            eq.add_seq_field(paragraph, "Equation", "Other", cached_text=5)
        self.assertEqual(paragraph._p.xml, before)

    def test_conversion_input_dependency_timeout_and_warning_failures_are_explicit(self):
        for value in ("", " ", None, "x" * 65537, r"\input{/tmp/source}"):
            with self.subTest(value=str(value)[:30]), self.assertRaises(ValueError):
                eq.latex_to_omml(value)
        with patch.object(eq.shutil, "which", return_value=None), self.assertRaisesRegex(RuntimeError, "Missing Pandoc"):
            eq.latex_to_omml("x^2")
        with patch.object(eq.subprocess, "run", side_effect=subprocess.TimeoutExpired("pandoc", 1)), self.assertRaisesRegex(RuntimeError, "timed out"):
            eq.latex_to_omml("x^2", pandoc="pandoc", timeout=1)
        for status, stderr in ((1, "parse error"), (0, "[WARNING] Could not convert TeX math")):
            with patch.object(eq.subprocess, "run", return_value=subprocess.CompletedProcess([], status, "", stderr)), \
                 self.assertRaisesRegex(RuntimeError, "conversion failed"):
                eq.latex_to_omml("x^2", pandoc="pandoc")
        with patch.object(eq.subprocess, "run", return_value=subprocess.CompletedProcess([], 0, "", "")), \
             self.assertRaisesRegex(RuntimeError, "without a DOCX"):
            eq.latex_to_omml("x^2", pandoc="pandoc")

    def test_converter_success_without_native_math_is_not_a_raster_fallback(self):
        for content in ('<w:p><w:r><w:t>raw LaTeX</w:t></w:r></w:p>',
                        '<w:p><w:r><w:drawing/></w:r></w:p>',
                        '<w:p><m:oMath><m:r><m:t>x</m:t></m:r></m:oMath><w:r><w:t>extra</w:t></w:r></w:p>'):
            def fake(command, **kwargs):
                output = Path(command[command.index("--output") + 1])
                with zipfile.ZipFile(output, "w") as package:
                    package.writestr("word/document.xml", f'<w:document xmlns:w="{W[1:-1]}" xmlns:m="{M[1:-1]}"><w:body>{content}</w:body></w:document>')
                return subprocess.CompletedProcess(command, 0, "", "")
            with self.subTest(content=content), patch.object(eq.subprocess, "run", side_effect=fake), self.assertRaises(RuntimeError):
                eq.latex_to_omml("x", pandoc="pandoc")

    def test_failed_numbered_conversion_does_not_partially_edit_document(self):
        document = Document()
        document.add_paragraph("未修改段落")
        before = document.element.xml
        with patch.object(eq, "latex_to_omml", side_effect=RuntimeError("conversion failure")), self.assertRaises(RuntimeError):
            eq.add_numbered_equation(document, "bad", bookmark="EquationBad")
        self.assertEqual(document.element.xml, before)

    @unittest.skipUnless(shutil.which("pandoc"), "Pandoc is not installed")
    def test_real_latex_fraction_sum_superscript_and_inline_are_native_omml(self):
        formula = r"\bar{x}=\frac{1}{n}\sum_{i=1}^{n}x_i^2"
        block = eq.latex_to_omml(formula)
        self.assertEqual(block.tag, M + "oMathPara")
        self.assertTrue(list(block.iter(M + "f")))
        self.assertTrue(list(block.iter(M + "nary")))
        self.assertTrue(list(block.iter(M + "sSup")) or list(block.iter(M + "sSubSup")))
        inline = eq.latex_to_omml("E=mc^2", display=False)
        self.assertEqual(inline.tag, M + "oMath")
        self.assertTrue(list(inline.iter(M + "sSup")))
        with self.assertRaises(RuntimeError):
            eq.latex_to_omml(r"\notARealMathCommand{x}")

    @unittest.skipUnless(shutil.which("pandoc"), "Pandoc is not installed")
    def test_numbered_equation_layout_source_retention_and_no_render_is_pending(self):
        document = Document()
        document.add_paragraph("保留原内容")
        original = self.root / "original.docx"
        document.save(original)
        baseline = hashlib.sha256(original.read_bytes()).hexdigest()
        working = Document(original)
        source = r"E=mc^2"
        result = eq.add_numbered_equation(working, source, bookmark="EquationEnergy")
        self.assertEqual(result["latex"], source)
        self.assertEqual(result["quality_status"], "pending")
        table = result["table"]
        self.assertEqual(len(table.columns), 3)
        self.assertTrue(list(table.cell(0, 1)._tc.iter(M + "oMath")))
        number = table.cell(0, 2).paragraphs[0]
        self.assertEqual(number.alignment, 2)
        self.assertEqual("".join(n.text or "" for n in number._p.iter(W + "t")), "(?)")
        self.assertTrue(all(edge.get(W + "val") == "nil" for edge in table._tbl.tblPr.find(W + "tblBorders")))
        self.assertEqual(table._tbl.tblPr.find(W + "tblCaption").get(W + "val"), "OfficeEquationLayout")
        working.save(self.root / "working.docx")
        self.assertEqual(hashlib.sha256(original.read_bytes()).hexdigest(), baseline)
        self.assertEqual(eq.validate_references(working)["native_equation_count"], 1)
        self.assertEqual(eq.validate_references(working)["quality_status"], "pending")

    def test_native_preservation_is_prefix_independent_and_detects_flattened_math(self):
        def package(name, prefix="m", content=None):
            output = self.root / name
            math = content if content is not None else f'<{prefix}:oMath><{prefix}:sSup><{prefix}:e><{prefix}:r><{prefix}:t>x</{prefix}:t></{prefix}:r></{prefix}:e><{prefix}:sup><{prefix}:r><{prefix}:t>2</{prefix}:t></{prefix}:r></{prefix}:sup></{prefix}:sSup></{prefix}:oMath>'
            with zipfile.ZipFile(output, "w") as target:
                target.writestr("[Content_Types].xml", "<Types/>")
                target.writestr("word/document.xml", f'<w:document xmlns:w="{W[1:-1]}" xmlns:{prefix}="{M[1:-1]}"><w:body><w:p>{math}</w:p></w:body></w:document>')
            return native.inspect_package(output)
        baseline = package("original.docx")
        same = package("different-prefix.docx", "eqn")
        self.assertEqual(native.preservation_errors(baseline, same), [])
        flat = package("flattened.docx", content='<w:r><w:t>x2</w:t></w:r>')
        self.assertTrue(any("math" in error for error in native.preservation_errors(baseline, flat)))
        wrong = package("wrong-structure.docx", content='<m:oMath><m:r><m:t>x2</m:t></m:r></m:oMath>')
        self.assertTrue(any("math" in error for error in native.preservation_errors(baseline, wrong)))

    def test_existing_global_flags_and_unrelated_or_locked_fields_are_preserved(self):
        for value in ("true", "false"):
            document = Document()
            setting = OxmlElement("w:updateFields")
            setting.set(W + "val", value)
            document.settings.element.append(setting)
            eq.add_seq_field(document.add_paragraph(), "Equation", "Energy")
            self.assertEqual(document.settings.element.find(W + "updateFields").get(W + "val"), value)
        for instruction, locked in (("ADDIN Zotero", False), ("CITATION source", False),
                                    ("BIBLIOGRAPHY", False), ("DDE external", False),
                                    ("DATE", False), ("SEQ Figure", True)):
            document = Document()
            existing = OxmlElement("w:fldSimple")
            existing.set(W + "instr", instruction)
            if locked:
                existing.set(W + "fldLock", "true")
            document.add_paragraph()._p.append(existing)
            before = etree.tostring(existing)
            eq.add_seq_field(document.add_paragraph(), "Equation", "Energy")
            with self.subTest(instruction=instruction, locked=locked):
                self.assertEqual(etree.tostring(existing), before)
                self.assertIsNone(document.settings.element.find(W + "updateFields"))
                self.assertEqual(eq.validate_references(document)["update_fields_on_open"], None)
                fields = list(document.element.iter(W + "fldSimple"))
                self.assertEqual(fields[-1].get(W + "dirty"), "true")

    def test_bookmarks_must_follow_start_order_and_part_but_may_overlap(self):
        document = Document()
        paragraph = document.add_paragraph()
        eq.add_seq_field(paragraph, "Equation", "Energy")
        start = next(paragraph._p.iter(W + "bookmarkStart"))
        end = next(paragraph._p.iter(W + "bookmarkEnd"))
        start.addprevious(end)
        self.assertEqual(eq.validate_references(document)["status"], "failed")
        with self.assertRaisesRegex(ValueError, "precedes"):
            eq.add_seq_field(document.add_paragraph(), "Equation", "Other")
        document = Document()
        p = document.add_paragraph()
        for tag, identifier, name in (("bookmarkStart", 0, "A"), ("bookmarkStart", 1, "B"),
                                      ("bookmarkEnd", 0, None), ("bookmarkEnd", 1, None)):
            node = OxmlElement("w:" + tag)
            node.set(W + "id", str(identifier))
            if name:
                node.set(W + "name", name)
            p._p.append(node)
        self.assertEqual(eq.validate_references(document)["status"], "references_valid")
        end = next(p._p.iter(W + "bookmarkEnd"))
        document.sections[0].header.paragraphs[0]._p.append(end)
        self.assertEqual(eq.validate_references(document)["status"], "failed")

    def test_quoted_reference_and_malformed_complex_fields_report_truthfully(self):
        document = Document()
        eq.add_seq_field(document.add_paragraph(), "Equation", "Energy")
        quoted = OxmlElement("w:fldSimple")
        quoted.set(W + "instr", ' REF "Energy" \\h ')
        document.add_paragraph()._p.append(quoted)
        self.assertEqual(eq.validate_references(document)["status"], "references_valid")
        for pattern in (("begin", "instruction"), ("end",), ("begin", "end"),
                        ("begin", "instruction", "separate"),
                        ("begin", "instruction", "separate", "separate", "end")):
            broken = Document()
            paragraph = broken.add_paragraph()
            for kind in pattern:
                if kind == "instruction":
                    item = OxmlElement("w:instrText")
                    item.text = " REF MissingTarget "
                else:
                    item = OxmlElement("w:fldChar")
                    item.set(W + "fldCharType", kind)
                run = OxmlElement("w:r")
                run.append(item)
                paragraph._p.append(run)
            result = eq.validate_references(broken)
            with self.subTest(pattern=pattern):
                self.assertEqual(result["status"], "failed", result)
                self.assertTrue(result["errors"])
                self.assertEqual(result["quality_status"], "pending")
        for instruction in ("REF", "SEQ", r"REF \h", 'REF "Missing'):
            broken = Document()
            field = OxmlElement("w:fldSimple")
            field.set(W + "instr", instruction)
            broken.add_paragraph()._p.append(field)
            self.assertEqual(eq.validate_references(broken)["status"], "failed")

    @unittest.skipUnless(shutil.which("pandoc"), "Pandoc is not installed")
    def test_consecutive_equation_layouts_have_a_real_paragraph_separator(self):
        document = Document()
        first = eq.add_numbered_equation(document, "E=mc^2", bookmark="Energy")
        second = eq.add_numbered_equation(document, r"x=\frac{a}{b}", bookmark="Fraction")
        between = first["table"]._tbl.getnext()
        self.assertEqual(between.tag, W + "p")
        self.assertIs(between.getnext(), second["table"]._tbl)
        self.assertFalse(list(between.iter(W + "vanish")))
        self.assertEqual(len(list(document.element.iter(M + "oMath"))), 2)
        self.assertEqual(eq.validate_references(document)["sequences"], {"Equation": 2})

    def test_native_math_guard_detects_changed_accent_operator_without_token_changes(self):
        document = Document()
        paragraph = document.add_paragraph()
        equation = OxmlElement("m:oMath")
        accent = OxmlElement("m:acc")
        props = OxmlElement("m:accPr")
        char = OxmlElement("m:chr")
        char.set(M + "val", "‾")
        props.append(char)
        accent.append(props)
        expression = OxmlElement("m:e")
        run, text = OxmlElement("m:r"), OxmlElement("m:t")
        text.text = "x"
        run.append(text)
        expression.append(run)
        accent.append(expression)
        equation.append(accent)
        paragraph._p.append(equation)
        before_path, after_path = self.root / "accent-before.docx", self.root / "accent-after.docx"
        document.save(before_path)
        char.set(M + "val", "´")
        document.save(after_path)
        before, after = native.inspect_package(before_path), native.inspect_package(after_path)
        self.assertEqual(before["text"], after["text"])
        self.assertTrue(any("math" in error for error in native.preservation_errors(before, after)))

    def test_math_topology_preserves_superscript_scope_and_ignores_run_splitting(self):
        def node(superscript):
            root = OxmlElement("m:oMath")
            exponent = OxmlElement("m:sSup")
            for role, value in (("e", "x"), ("sup", superscript)):
                child = OxmlElement("m:" + role)
                run, text = OxmlElement("m:r"), OxmlElement("m:t")
                text.text = value
                run.append(text)
                child.append(run)
                exponent.append(child)
            root.append(exponent)
            return root
        first, second = node("2"), node("2+y")
        run, text = OxmlElement("m:r"), OxmlElement("m:t")
        text.text = "+y"
        run.append(text)
        first.append(run)
        self.assertNotEqual(native._math_topology(first), native._math_topology(second))
        split = deepcopy(first)
        tail = split[-1]
        tail[0].text = "+"
        other, other_text = OxmlElement("m:r"), OxmlElement("m:t")
        other_text.text = "y"
        other.append(other_text)
        split.append(other)
        self.assertEqual(native._math_topology(first), native._math_topology(split))

    def test_generic_footnotes_part_participates_in_refs_ids_math_and_update_scope(self):
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        document = Document()
        blob = f'''<w:footnotes xmlns:w="{W[1:-1]}" xmlns:m="{M[1:-1]}"><w:footnote w:id="1"><w:p>
        <w:bookmarkStart w:id="7" w:name="FootnoteNumber"/><w:fldSimple w:instr="SEQ Equation"/><w:bookmarkEnd w:id="7"/>
        <w:fldSimple w:instr="REF MissingTarget"/><w:fldSimple w:instr="CITATION RealSource"/>
        <m:oMath><m:r><m:t>x</m:t></m:r></m:oMath></w:p></w:footnote></w:footnotes>'''.encode()
        part = Part(PackURI("/word/footnotes.xml"),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                    blob=blob, package=document.part.package)
        document.part.relate_to(part, RT.FOOTNOTES)
        eq.add_seq_field(document.add_paragraph(), "Equation", "MainNumber")
        report = eq.validate_references(document)
        self.assertEqual(report["status"], "failed", report)
        self.assertIn({"command": "REF", "bookmark": "MissingTarget"}, report["references"])
        self.assertIn("FootnoteNumber", report["bookmark_names"])
        self.assertEqual(report["sequences"], {"Equation": 2})
        self.assertEqual(report["native_equation_count"], 1)
        self.assertIsNone(document.settings.element.find(W + "updateFields"))
        main = next(document.element.iter(W + "bookmarkStart"))
        self.assertEqual(main.get(W + "id"), "8")
        self.assertEqual(part.blob, blob)

    def test_word_macro_orders_seq_ref_and_pagination_without_managed_fields_or_autorun(self):
        macro = (REPO / "manifests/office/docx/macros/OfficeLayout.bas").read_text()
        start = macro.index("Private Sub RefreshNativeFields")
        body = macro[start:macro.index("Private Function UpdateFieldPass", start)]
        one, two, three = (body.index(f"UpdateFieldPass(doc, {n})") for n in (1, 2, 3))
        self.assertLess(body.index("UpdateFieldPass(doc, 0)"), one)
        self.assertLess(one, two)
        self.assertLess(two, body.index("doc.Repaginate"))
        self.assertLess(body.index("doc.Repaginate"), three)
        self.assertIn("item.Type = wdFieldSequence", macro)
        self.assertIn("If Not item.Locked And IsInternalField(item)", macro)
        self.assertIn("wdFieldAddin, wdFieldCitation, wdFieldBibliography", macro)
        executable = "\n".join(line for line in macro.splitlines() if not line.lstrip().startswith("'"))
        for forbidden in ("Sub AutoOpen", "Sub Document_Open", ".Save", "SaveAs", "VBProject", "TrackRevisions =", "AcceptAllRevisions"):
            self.assertNotIn(forbidden, executable)


if __name__ == "__main__":
    unittest.main()
