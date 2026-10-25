"""Read-only macro evidence and simulated, explicitly scoped chart decisions."""
import hashlib
import importlib.util
import os
from pathlib import Path
import re
import shutil
import tempfile
import unittest
import warnings
import xml.etree.ElementTree as ET
import zipfile

REPO = Path(__file__).resolve().parents[3]
SCRIPTS = REPO / "manifests/office/xlsx/scripts"
spec = importlib.util.spec_from_file_location("office_workbook_features_test", SCRIPTS / "workbook_features.py")
features = importlib.util.module_from_spec(spec)
spec.loader.exec_module(features)

try:
    import openpyxl
    import xlsxwriter
except ImportError:
    openpyxl = xlsxwriter = None

FIXTURE = Path(os.environ["OFFICE_VBA_FIXTURE"]) if os.environ.get("OFFICE_VBA_FIXTURE") else None
FIXTURE_SHA = "0ced1464b3677e98f5e3a8c5d80135e18dc98dca39299f1a8cfd2a00999fbf9f"
# Official XlsxWriter example at pinned commit 227043ddc790c14887245012f85471c3a7810354.
# This binary is NOT OfficeWorkbook.bas compiled, and no test executes it.


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def rewrite(source, output, change):
    with zipfile.ZipFile(source) as incoming:
        files = {name: incoming.read(name) for name in incoming.namelist()}
    change(files)
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as outgoing:
        for name, data in files.items():
            outgoing.writestr(name, data)


def inventory_fixture(path, *, macro=False, controls=False):
    """Synthetic inert OOXML fixture for package inspection, not Excel execution."""
    ns, rel, ct = features.S, features.P, features.C
    main_type = features.MAIN_TYPES[path.suffix]
    files = {
        "[Content_Types].xml": f'<Types xmlns="{ct}"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="{main_type}"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>',
        "_rels/.rels": f'<Relationships xmlns="{rel}"><Relationship Id="rId1" Type="{features.R}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook xmlns="{ns}" xmlns:r="{features.R}"><workbookPr codeName="ThisWorkbook"/><sheets><sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": f'<Relationships xmlns="{rel}"><Relationship Id="rId1" Type="{features.R}/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": f'<worksheet xmlns="{ns}" xmlns:r="{features.R}"><sheetPr codeName="Sheet1"/><sheetData/></worksheet>',
    }
    def add_type(part, kind):
        files["[Content_Types].xml"] = files["[Content_Types].xml"].replace("</Types>", f'<Override PartName="/{part}" ContentType="{kind}"/></Types>')
    def add_rel(source, target, kind, rid):
        files[source] = files[source].replace("</Relationships>", f'<Relationship Id="{rid}" Type="{kind}" Target="{target}"/></Relationships>')
    if macro:
        files["xl/vbaProject.bin"] = b"SYNTHETIC INERT INVENTORY DATA; NOT EXECUTABLE VBA"
        add_type("xl/vbaProject.bin", "application/vnd.ms-office.vbaProject")
        add_rel("xl/_rels/workbook.xml.rels", "vbaProject.bin", "http://schemas.microsoft.com/office/2006/relationships/vbaProject", "rId2")
    if controls:
        files["xl/ctrlProps/ctrlProp1.xml"] = '<formControlPr xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" objectType="Button"/>'
        files["xl/worksheets/_rels/sheet1.xml.rels"] = f'<Relationships xmlns="{rel}"/>'
        add_type("xl/ctrlProps/ctrlProp1.xml", "application/vnd.ms-excel.controlproperties+xml")
        add_rel("xl/worksheets/_rels/sheet1.xml.rels", "../ctrlProps/ctrlProp1.xml", "http://schemas.microsoft.com/office/2006/relationships/ctrlProp", "rId3")
        files["xl/worksheets/sheet1.xml"] = files["xl/worksheets/sheet1.xml"].replace('</worksheet>', '<controls><control shapeId="1025" name="Button 1" r:id="rId3"/></controls></worksheet>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        for name, content in files.items():
            package.writestr(name, content)


class WorkbookFeatureTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="office workbook features ")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_decisions_pending_yes_no_and_evidence_scope(self):
        pending = features.chart_decision()
        self.assertEqual(pending["decision"], "pending")
        with self.assertRaises(features.ChartConfirmationRequired):
            features.require_chart_confirmation(pending)
        for answer, allowed in (("yes", True), ("no", False)):
            decision = features.chart_decision(answer, source="SIMULATED user answer for test", scope="sales-summary")
            self.assertIs(features.require_chart_confirmation(decision, scope="sales-summary"), allowed)
            with self.assertRaises(features.ChartConfirmationRequired):
                features.require_chart_confirmation(decision, scope="another-chart")
        for answer in (True, "YES", "maybe", 1):
            with self.assertRaises(ValueError):
                features.chart_decision(answer, source="SIMULATED")
        with self.assertRaises(ValueError):
            features.chart_decision("yes")
        self.assertTrue(features.require_chart_confirmation(creating_new_charts=False))

    def test_macro_approval_is_separate_and_operation_scoped(self):
        self.assertTrue(features.require_macro_confirmation(mode="preserve"))
        for mode in ("create", "modify", "invoke"):
            with self.subTest(mode=mode), self.assertRaises(features.MacroConfirmationRequired):
                features.require_macro_confirmation(mode=mode)
        decision = features.macro_decision("yes", source="SIMULATED user approval", requested_modes=("create", "modify"))
        self.assertTrue(features.require_macro_confirmation(decision, mode="create"))
        with self.assertRaises(features.MacroConfirmationRequired):
            features.require_macro_confirmation(decision, mode="invoke")
        with self.assertRaises(features.MacroConfirmationRequired):
            features.require_macro_confirmation(decision, mode="create", scope="another-workbook")
        denied = features.macro_decision("no", source="SIMULATED refusal", requested_modes=["invoke"])
        self.assertFalse(features.require_macro_confirmation(denied, mode="invoke"))
        with self.assertRaises(features.MacroConfirmationRequired):
            features.require_macro_confirmation(features.chart_decision("yes", source="SIMULATED"), mode="invoke")
        for modes in ("invoke", [], ["all"], [None]):
            with self.subTest(modes=modes), self.assertRaises(ValueError):
                features.macro_decision(requested_modes=modes)

    def test_four_package_types_source_hash_and_static_pending(self):
        for extension in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            with self.subTest(extension=extension):
                source = self.root / ("中文 source" + extension)
                inventory_fixture(source, macro=extension.endswith("m"))
                before = sha(source)
                result = features.inspect_workbook(source)
                self.assertEqual(result["status"], "inspected", result)
                self.assertEqual(result["source_sha256"], before)
                self.assertEqual(sha(source), before)
                self.assertTrue(result["source_unchanged"])
                self.assertEqual(result["quality_status"], "pending")
                self.assertEqual(result["macro_execution_status"], "not_run")

    def test_macro_bytes_type_relationship_codename_and_control_binding_changes_fail(self):
        source = self.root / "source.xlsm"
        inventory_fixture(source, macro=True, controls=True)
        mutations = {
            "binary": lambda files: files.__setitem__("xl/vbaProject.bin", b"changed inert bytes"),
            "type": lambda files: files.__setitem__("[Content_Types].xml", files["[Content_Types].xml"].replace(b"application/vnd.ms-office.vbaProject", b"application/octet-stream")),
            "rel": lambda files: files.__setitem__("xl/_rels/workbook.xml.rels", files["xl/_rels/workbook.xml.rels"].replace(b"relationships/vbaProject", b"relationships/worksheet")),
            "codeName": lambda files: files.__setitem__("xl/worksheets/sheet1.xml", files["xl/worksheets/sheet1.xml"].replace(b'codeName="Sheet1"', b'codeName="Changed"')),
            "control": lambda files: files.__setitem__("xl/worksheets/sheet1.xml", files["xl/worksheets/sheet1.xml"].replace(b'shapeId="1025"', b'shapeId="9999"')),
            "controlpart": lambda files: files.__setitem__("xl/ctrlProps/ctrlProp1.xml", files["xl/ctrlProps/ctrlProp1.xml"].replace(b'Button', b'CheckBox')),
        }
        for label, mutation in mutations.items():
            with self.subTest(label=label):
                output = self.root / (label + ".xlsm")
                rewrite(source, output, mutation)
                result = features.verify_macro_preservation(source, output)
                self.assertFalse(result["preservation_passed"], result)
                self.assertTrue(result["errors"])
                self.assertEqual(result["macro_execution_status"], "not_run")

    def test_renamed_vba_project_is_protected_by_relationship_and_content_type(self):
        source, renamed, lost = [self.root / name for name in ("source.xlsm", "renamed.xlsm", "lost.xlsm")]
        inventory_fixture(source, macro=True)
        def rename(files):
            files["xl/macros.bin"] = files.pop("xl/vbaProject.bin")
            for part in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                files[part] = files[part].replace(b"vbaProject.bin", b"macros.bin")
        rewrite(source, renamed, rename)
        self.assertEqual(features.inspect_workbook(renamed)["vba_projects"], ["xl/macros.bin"])
        def remove(files):
            del files["xl/macros.bin"]
            for part in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                root = ET.fromstring(files[part])
                for node in list(root):
                    if "macros.bin" in " ".join(node.attrib.values()):
                        root.remove(node)
                files[part] = ET.tostring(root)
        rewrite(renamed, lost, remove)
        self.assertFalse(features.verify_macro_preservation(renamed, lost)["preservation_passed"])

    def test_customui_signature_activex_vml_xlm_and_dependency_graph_are_inventoried(self):
        source = self.root / "source.xlsm"
        inventory_fixture(source, macro=True)
        def enrich(files):
            files.update({"customUI/customUI.xml": b'<customUI/>', "customUI/_rels/customUI.xml.rels": f'<Relationships xmlns="{features.P}"><Relationship Id="img1" Type="{features.R}/image" Target="../xl/media/support.png"/></Relationships>'.encode(),
                          "xl/media/support.png": b"INERT image fixture", "xl/vbaProjectSignature.bin": b"INERT signature fixture",
                          "xl/activeX/activeX1.bin": b"INERT ActiveX fixture", "xl/drawings/vmlDrawing1.vml": b'<xml/>',
                          "xl/macrosheets/sheet1.xml": b'<macrosheet/>'})
        enriched = self.root / "enriched.xlsm"
        rewrite(source, enriched, enrich)
        inventory = features.inspect_workbook(enriched)
        for part in ("customUI/customUI.xml", "xl/media/support.png", "xl/vbaProjectSignature.bin", "xl/activeX/activeX1.bin", "xl/drawings/vmlDrawing1.vml", "xl/macrosheets/sheet1.xml"):
            self.assertIn(part, inventory["macro_related_parts"])
        changed = self.root / "changed.xlsm"
        rewrite(enriched, changed, lambda files: files.__setitem__("xl/media/support.png", b"changed dependency"))
        self.assertFalse(features.verify_macro_preservation(enriched, changed)["preservation_passed"])

    def test_chart_auxiliary_parts_are_not_charts(self):
        source = self.root / "source.xlsx"
        inventory_fixture(source)
        output = self.root / "chart.xlsx"
        def enrich(files):
            files.update({"xl/charts/chart1.xml": b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart/></c:chartSpace>',
                          "xl/charts/style1.xml": b'<chartStyle/>', "xl/charts/colors1.xml": b'<colorStyle/>', "xl/charts/drawing1.xml": b'<drawing/>'})
        rewrite(source, output, enrich)
        result = features.inspect_workbook(output)
        self.assertEqual(result["native_chart_count"], 1)
        self.assertEqual(len(result["chart_auxiliary_parts"]), 3)

    def test_drawing_macro_and_xlm_defined_name_bindings_survive_or_fail(self):
        source = self.root / "source.xlsm"
        inventory_fixture(source, macro=True)
        bound = self.root / "bound.xlsm"
        def enrich(files):
            files["xl/workbook.xml"] = files["xl/workbook.xml"].replace(b'</workbook>', b'<definedNames><definedName name="_xlnm.Auto_Open" localSheetId="0">MacroSheet!$A$1</definedName><definedName name="CustomMacro" vbProcedure="1">Data!$A$2</definedName><definedName name="OrdinaryArea">Data!$A$3</definedName></definedNames></workbook>')
            files["xl/drawings/drawing1.xml"] = b'<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"><xdr:sp macro="[0]!Module1.ReviewedAction"><xdr:nvSpPr><xdr:cNvPr id="2" name="Macro Shape"/></xdr:nvSpPr></xdr:sp></xdr:wsDr>'
        rewrite(source, bound, enrich)
        before = features.inspect_workbook(bound)
        self.assertEqual(len(before["macro_bindings"]["xl/workbook.xml"]), 2)
        self.assertEqual(len(before["macro_bindings"]["xl/drawings/drawing1.xml"]), 1)
        for label, part, old, new in (
                ("name", "xl/workbook.xml", b'MacroSheet!$A$1', b'MacroSheet!$A$5'),
                ("drawing", "xl/drawings/drawing1.xml", b' macro="[0]!Module1.ReviewedAction"', b'')):
            with self.subTest(label=label):
                changed = self.root / (label + ".xlsm")
                rewrite(bound, changed, lambda files: files.__setitem__(part, files[part].replace(old, new)))
                result = features.verify_macro_preservation(bound, changed)
                self.assertIn("Protected feature changed: macro_bindings", result["errors"])
        # Adding a regular chart/drawing does not invalidate retained bindings.
        chart = self.root / "chart-added.xlsm"
        rewrite(bound, chart, lambda files: files.__setitem__('xl/charts/chart1.xml', b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart/></c:chartSpace>'))
        self.assertTrue(features.verify_macro_preservation(bound, chart)["preservation_passed"])

    def test_mismatch_missing_relationship_dtd_and_same_path_fail(self):
        source = self.root / "source.xlsm"
        inventory_fixture(source, macro=True)
        renamed = self.root / "wrong.xlsx"
        shutil.copyfile(source, renamed)
        self.assertEqual(features.inspect_workbook(renamed)["status"], "failed")
        with self.assertRaises(ValueError):
            features.verify_macro_preservation(source, source)
        output = self.root / "missing.xlsm"
        rewrite(source, output, lambda files: files.pop("xl/vbaProject.bin"))
        self.assertEqual(features.inspect_workbook(output)["status"], "failed")
        output = self.root / "entity.xlsm"
        rewrite(source, output, lambda files: files.__setitem__("xl/workbook.xml", b'<!DOCTYPE workbook [<!ENTITY x "y">]><workbook/>'))
        with self.assertRaises(ValueError):
            features.inspect_workbook(output)

    def test_duplicate_member_and_escaping_target_fail_without_source_change(self):
        source = self.root / "source.xlsx"
        inventory_fixture(source)
        original = sha(source)
        duplicate = self.root / "duplicate.xlsx"
        shutil.copyfile(source, duplicate)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(duplicate, "a") as z:
                z.writestr("xl/workbook.xml", b"<workbook/>")
        with self.assertRaises(ValueError):
            features.inspect_workbook(duplicate)
        escaped = self.root / "escaped.xlsx"
        rewrite(source, escaped, lambda files: files.__setitem__("xl/_rels/workbook.xml.rels", files["xl/_rels/workbook.xml.rels"].replace(b"worksheets/sheet1.xml", b"../../outside.xml")))
        with self.assertRaises(ValueError):
            features.inspect_workbook(escaped)
        self.assertEqual(sha(source), original)

    def test_vba_has_explicit_inputs_confirmation_and_no_automatic_execution(self):
        path = REPO / "manifests/office/xlsx/macros/OfficeWorkbook.bas"
        text = path.read_text()
        code = "\n".join(line for line in text.splitlines() if not line.lstrip().startswith("'"))
        self.assertIn('Attribute VB_Name = "OfficeWorkbook"', code)
        for name in ("CalculateOfficeRange", "ConfigureOfficePrintArea", "AddOfficeNativeChart"):
            self.assertRegex(code, rf"Public (?:Sub|Function) {name}\(ByVal wb As Workbook")
        for forbidden in (r"\bAuto_Open\b", r"\bWorkbook_Open\b", r"\bActiveWorkbook\b", r"\bActiveSheet\b", r"\.Save(?:As)?\b", r"VBProject", r"AutomationSecurity", r"Application\.Run", r"\.RefreshAll\b", r"Application\.Calculate\b"):
            self.assertNotRegex(code, forbidden)
        body = code[code.index("Public Function AddOfficeNativeChart"):]
        self.assertLess(body.index("RequireOfficeApproval macroConfirmed"), body.index("ChartObjects.Add"))
        self.assertLess(body.index("If Not chartConfirmed Then Exit Function"), body.index("ChartObjects.Add"))
        self.assertIn("Optional ByVal macroConfirmed As Boolean = False", body)
        self.assertIn("wb.Application.EnableEvents = previousEvents", body)


@unittest.skipUnless(openpyxl and xlsxwriter, "Optional openpyxl/XlsxWriter not installed")
class ActualWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="office actual workbook ")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_actual_chart_yes_no_pending_are_simulated_and_do_not_modify_existing_charts(self):
        from openpyxl.chart import BarChart, Reference
        for answer in ("yes", "no", "pending"):
            with self.subTest(answer=answer):
                book = openpyxl.Workbook()
                sheet = book.active
                sheet.append(["Month", "Sales"])
                sheet.append(["Jan", 4])
                sheet.append(["Feb", 6])
                decision = features.chart_decision(answer, source="SIMULATED test reply" if answer != "pending" else None)
                try:
                    allowed = features.require_chart_confirmation(decision)
                except features.ChartConfirmationRequired:
                    allowed = False
                if allowed:
                    chart = BarChart()
                    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
                    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
                    sheet.add_chart(chart, "D2")
                # Pending/no does not block unrelated data work or invent charts.
                sheet["B4"] = "=SUM(B2:B3)"
                output = self.root / (answer + ".xlsx")
                book.save(output)
                book.close()
                self.assertEqual(features.inspect_workbook(output)["native_chart_count"], int(answer == "yes"))
        yes = self.root / "yes.xlsx"
        working = openpyxl.load_workbook(yes)
        self.assertTrue(features.require_chart_confirmation(creating_new_charts=False))
        working.active["A5"] = "unrelated edit"
        working.save(self.root / "existing.xlsx")
        working.close()
        self.assertEqual(features.inspect_workbook(self.root / "existing.xlsx")["native_chart_count"], 1)

    @unittest.skipUnless(FIXTURE is not None and FIXTURE.is_file(),
                         "Pinned official VBA fixture unavailable; set OFFICE_VBA_FIXTURE")
    def test_actual_official_vba_keep_vba_copy_and_wrongsave_failure(self):
        self.assertEqual(sha(FIXTURE), FIXTURE_SHA, "Only the pinned official test binary is accepted")
        source, output, lost = [self.root / name for name in ("source.xlsm", "working.xlsm", "lost.xlsm")]
        book = xlsxwriter.Workbook(source)
        self.assertEqual(book.add_vba_project(str(FIXTURE)), 0)
        book.set_vba_name("ThisWorkbook")
        sheet = book.add_worksheet("Data")
        sheet.set_vba_name("Sheet1")
        sheet.write_row("A1", ["Month", "Sales"])
        sheet.write_row("A2", ["Jan", 4])
        book.close()
        original = sha(source)
        working = openpyxl.load_workbook(source, keep_vba=True, data_only=False)
        working["Data"]["B2"] = 8
        working.save(output)
        working.close()
        evidence = features.verify_macro_preservation(source, output)
        self.assertEqual(evidence["status"], "preserved", evidence["errors"])
        self.assertEqual(evidence["quality_status"], "pending")
        self.assertEqual(evidence["macro_execution_status"], "not_run")
        self.assertEqual(sha(source), original)
        wrong = openpyxl.load_workbook(source, keep_vba=False, data_only=False)
        wrong.save(lost)
        wrong.close()
        failure = features.verify_macro_preservation(source, lost)
        self.assertFalse(failure["preservation_passed"])
        self.assertTrue(any("macro_related_parts" in item for item in failure["errors"]))
        self.assertEqual(sha(source), original)


if __name__ == "__main__":
    unittest.main()
