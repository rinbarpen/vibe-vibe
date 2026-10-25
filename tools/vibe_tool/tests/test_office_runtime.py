"""Runtime guards use mocks; real format checks run when their libraries exist."""

import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch


REPO_ROOT = Path(__file__).resolve().parents[3]
SCRIPTS = REPO_ROOT / "manifests/office/scripts"
RECALC_SCRIPT = REPO_ROOT / "manifests/office/xlsx/scripts/recalc.py"


def load_script(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    # render.py intentionally imports its installed sibling verify.py.
    sys.path.insert(0, str(path.parent))
    try:
        spec.loader.exec_module(module)
    finally:
        sys.path.pop(0)
    return module


doctor = load_script("office_test_doctor", SCRIPTS / "doctor.py")
verify = load_script("office_test_verify", SCRIPTS / "verify.py")
render = load_script("office_test_render", SCRIPTS / "render.py")
recalc = load_script("office_test_recalc", RECALC_SCRIPT)


class RuntimeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="office runtime ")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def source(self, name="input.docx", data=b"unchanged fixture"):
        path = self.root / name
        path.write_bytes(data)
        return path

    def cli(self, script, *args, env=None, timeout=30):
        return subprocess.run([sys.executable, str(script), *map(str, args)], cwd=self.root,
                              env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1", **(env or {})},
                              capture_output=True, text=True, timeout=timeout)

    def test_doctor_missing_tools_modules_and_fonts_is_serializable(self):
        with patch.object(doctor.shutil, "which", return_value=None), \
             patch.object(doctor.importlib.util, "find_spec", return_value=None), \
             patch.object(doctor.subprocess, "run") as run:
            report = doctor.inspect()
        run.assert_not_called()
        self.assertEqual(json.loads(json.dumps(report))["status"], "dependencies_missing")
        self.assertIn("LibreOffice", report["missing"])
        self.assertIn("CJK fonts (or fontconfig evidence)", report["missing"])
        self.assertIn("python-docx", report["missing"])
        self.assertIn("node:pptxgenjs", report["missing"])
        self.assertEqual(list(self.root.iterdir()), [])

    def test_doctor_missing_fonts_is_not_ready_even_with_all_required_tools(self):
        available = {"node", "soffice", "pdfinfo", "pdftoppm", "fc-list"}

        def run(command, **kwargs):
            return subprocess.CompletedProcess(command, 0, "" if command[0] == "fc-list" else "module", "")

        with patch.object(doctor.shutil, "which", side_effect=lambda x: x if x in available else None), \
             patch.object(doctor.importlib.util, "find_spec", return_value=object()), \
             patch.object(doctor.subprocess, "run", side_effect=run):
            report = doctor.inspect()
        self.assertEqual(report["status"], "dependencies_missing")
        self.assertEqual(report["missing"], ["CJK fonts (or fontconfig evidence)"])
        self.assertNotIn("tesseract", report["missing"])

    def test_doctor_cli_missing_path_returns_json_nonzero_without_writes(self):
        before = list(self.root.iterdir())
        result = self.cli(SCRIPTS / "doctor.py", env={"PATH": ""})
        self.assertEqual(result.returncode, 1, result.stderr)
        report = json.loads(result.stdout)
        self.assertEqual(report["status"], "dependencies_missing")
        self.assertIn("LibreOffice", report["missing"])
        self.assertEqual(list(self.root.iterdir()), before)

    def test_verify_corrupt_packages_report_failure_and_do_not_mutate(self):
        for fmt in ("docx", "xlsx", "pptx", "pdf"):
            with self.subTest(fmt=fmt):
                source = self.source(f"broken.{fmt}", b"broken file")
                before = source.read_bytes()
                report = verify.verify(source)
                self.assertEqual(report["status"], "failed")
                self.assertEqual(report["quality_status"], "pending")
                self.assertTrue(report["errors"])
                self.assertTrue(report["checks"]["input_unchanged"])
                self.assertEqual(source.read_bytes(), before)

    def test_recalc_rejected_macro_format_records_actual_unchanged_source(self):
        source = self.source("input.xlsm", b"rejected before opening package")
        output = self.root / "output.xlsm"
        with patch.object(recalc.subprocess, "run") as run:
            report = recalc.recalculate(source, 30, output)
        run.assert_not_called()
        self.assertEqual(report["status"], "failed")
        self.assertEqual(report["source_sha256_before"], recalc.digest(source))
        self.assertEqual(report["source_sha256_after"], recalc.digest(source))
        self.assertTrue(report["source_unchanged"])
        self.assertFalse(output.exists())

    def test_verify_report_collision_with_source_or_baseline_is_rejected(self):
        source = self.source()
        baseline = self.source("baseline.docx", b"baseline")
        before = source.read_bytes(), baseline.read_bytes()
        for output in (source, baseline):
            result = self.cli(SCRIPTS / "verify.py", source, "--baseline", baseline, "--output", output)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("Report path must differ", result.stderr)
        self.assertEqual((source.read_bytes(), baseline.read_bytes()), before)

    def test_verify_corruption_writes_truthful_json_and_baseline_hash(self):
        source = self.source()
        baseline = self.source("baseline.docx", b"baseline")
        output = self.root / "reports/verification.json"
        result = self.cli(SCRIPTS / "verify.py", source, "--baseline", baseline, "--output", output)
        self.assertEqual(result.returncode, 1, result.stderr)
        report = json.loads(output.read_text())
        self.assertEqual(report["baseline"]["sha256"], verify.sha256(baseline))
        self.assertEqual(report["sha256"], verify.sha256(source))
        self.assertEqual(report["status"], "failed")

    def test_verify_and_render_reject_unsupported_format_without_source_changes(self):
        source = self.source("input.csv")
        before = source.read_bytes()
        report = verify.verify(source)
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any("Expected .docx" in x for x in report["errors"]))
        out = self.root / "rendered"
        with self.assertRaisesRegex(ValueError, "Expected .docx"):
            render.render(source, out)
        self.assertFalse(out.exists())
        self.assertEqual(source.read_bytes(), before)

    def test_render_missing_office_is_pending_and_preserves_source(self):
        source = self.source()
        out = self.root / "rendered"
        with patch.object(render.shutil, "which", return_value=None):
            report = render.render(source, out)
        self.assertEqual(report["status"], "failed")
        self.assertEqual(report["quality_status"], "pending")
        self.assertIn("Missing LibreOffice", report["error"])
        self.assertTrue(report["input_unchanged"])
        self.assertEqual(json.loads((out / "render.json").read_text()), report)

    def test_render_zero_exit_without_pdf_is_failure(self):
        source = self.source()
        with patch.object(render.shutil, "which", return_value="/fixture/libreoffice"), \
             patch.object(render.subprocess, "run", return_value=subprocess.CompletedProcess([], 0, "success", "")):
            report = render.render(source, self.root / "rendered")
        self.assertEqual(report["status"], "failed")
        self.assertIn("produced no PDF", report["error"])
        self.assertEqual(report["commands"][0]["exit_status"], 0)
        self.assertTrue(report["input_unchanged"])
        command = report["commands"][0]["command"]
        self.assertTrue(any(x.startswith("-env:UserInstallation=file:") for x in command))
        self.assertNotEqual(Path(command[-1]), source)

    def test_render_timeout_is_recorded_as_timeout_not_success(self):
        source = self.source()
        with patch.object(render.shutil, "which", return_value="/fixture/libreoffice"), \
             patch.object(render.subprocess, "run", side_effect=subprocess.TimeoutExpired("libreoffice", 2, output="partial")):
            report = render.render(source, self.root / "rendered", timeout=2)
        self.assertEqual(report["status"], "failed")
        self.assertIn("timed out", report["error"])
        self.assertTrue(report["commands"][0]["timed_out"])
        self.assertIsNone(report["commands"][0]["exit_status"])
        self.assertTrue(report["input_unchanged"])

    def test_render_existing_evidence_directory_is_never_reused(self):
        source = self.source()
        out = self.root / "rendered"
        out.mkdir()
        sentinel = out / "previous evidence.txt"
        sentinel.write_text("keep")
        with self.assertRaisesRegex(ValueError, "must be empty"):
            render.render(source, out)
        self.assertEqual(sentinel.read_text(), "keep")

    @staticmethod
    def workbook_inspection(errors=None, missing=None):
        return {"formulas": {"Sheet1!B1": "=A1*2"}, "formula_count": 1,
                "errors": errors or [], "missing_caches": missing or [],
                "sheet_names": ["Sheet1"], "charts": 1}

    def fake_conversion(self, command, **kwargs):
        incoming = Path(command[-1])
        outgoing = Path(command[command.index("--outdir") + 1])
        (outgoing / incoming.name).write_bytes(b"recalculated workbook")
        return subprocess.CompletedProcess(command, 0, "converted", "")

    def test_recalc_fake_office_zero_exit_without_new_output_retains_source(self):
        source = self.source("working.xlsx")
        with patch.object(recalc, "inspect", return_value=self.workbook_inspection()), \
             patch.object(recalc, "prepare_recalculation", side_effect=shutil.copy2), \
             patch.object(recalc.shutil, "which", return_value="/fixture/libreoffice"), \
             patch.object(recalc.subprocess, "run", return_value=subprocess.CompletedProcess([], 0, "", "")):
            report = recalc.recalculate(source, 5)
        self.assertEqual(report["status"], "failed")
        self.assertIn("without creating", report["error"])
        self.assertTrue(report["source_unchanged"])
        self.assertEqual(source.read_bytes(), b"unchanged fixture")

    def test_recalc_missing_office_retains_working_copy(self):
        source = self.source("working.xlsx")
        with patch.object(recalc, "inspect", return_value=self.workbook_inspection()), \
             patch.object(recalc.shutil, "which", return_value=None), \
             patch.object(recalc.subprocess, "run") as run:
            report = recalc.recalculate(source, 5)
        run.assert_not_called()
        self.assertEqual(report["status"], "failed")
        self.assertIn("Missing LibreOffice", report["error"])
        self.assertTrue(report["source_unchanged"])

    def test_recalc_rejects_non_xlsx_and_nonpositive_timeout_without_execution(self):
        source = self.source("working.xlsx")
        csv = self.source("working.csv")
        for file, timeout, output in ((source, 0, None), (csv, 5, None),
                                      (source, 5, self.root / "result.csv")):
            with self.subTest(file=file, timeout=timeout, output=output), \
                 patch.object(recalc.subprocess, "run") as run:
                report = recalc.recalculate(file, timeout, output)
            run.assert_not_called()
            self.assertEqual(report["status"], "failed")
            self.assertIn("positive timeout", report["error"])
            self.assertEqual(file.read_bytes(), b"unchanged fixture")

    def test_recalc_timeout_and_nonzero_exit_retain_input_and_existing_output(self):
        source = self.source("working.xlsx")
        output = self.source("result.xlsx", b"old result")
        for failure in (subprocess.TimeoutExpired("libreoffice", 2), subprocess.CompletedProcess([], 7, "", "failed")):
            with self.subTest(failure=type(failure).__name__), \
                 patch.object(recalc, "inspect", return_value=self.workbook_inspection()), \
                 patch.object(recalc, "prepare_recalculation", side_effect=shutil.copy2), \
                 patch.object(recalc.shutil, "which", return_value="/fixture/libreoffice"), \
                 patch.object(recalc.subprocess, "run") as run:
                if isinstance(failure, Exception):
                    run.side_effect = failure
                else:
                    run.return_value = failure
                report = recalc.recalculate(source, 2, output)
            self.assertEqual(report["status"], "failed")
            self.assertTrue(report["source_unchanged"])
            self.assertEqual(output.read_bytes(), b"old result")
            if isinstance(failure, Exception):
                self.assertTrue(report["commands"][0]["timed_out"])
            else:
                self.assertEqual(report["commands"][0]["exit_status"], 7)

    def test_recalc_formula_errors_absent_caches_or_changed_formulas_retain_source(self):
        source = self.source("working.xlsx")
        variants = [self.workbook_inspection(errors=[{"cell": "Sheet1!B1", "error": "#DIV/0!"}]),
                    self.workbook_inspection(missing=["Sheet1!B1"]),
                    {**self.workbook_inspection(), "formulas": {"Sheet1!B1": "=A1*3"}}]
        for after in variants:
            with self.subTest(after=after), \
                 patch.object(recalc, "inspect", side_effect=[self.workbook_inspection(), after]), \
                 patch.object(recalc, "prepare_recalculation", side_effect=shutil.copy2), \
                 patch.object(recalc.shutil, "which", return_value="/fixture/libreoffice"), \
                 patch.object(recalc.subprocess, "run", side_effect=self.fake_conversion):
                report = recalc.recalculate(source, 5)
            self.assertEqual(report["status"], "failed")
            self.assertTrue(report["source_unchanged"])
            self.assertEqual(source.read_bytes(), b"unchanged fixture")

    def test_recalc_success_uses_isolated_input_output_profile_and_preserves_source(self):
        source = self.source("working.xlsx")
        output = self.root / "results/recalculated.xlsx"
        with patch.object(recalc, "inspect", return_value=self.workbook_inspection()), \
             patch.object(recalc, "prepare_recalculation", side_effect=shutil.copy2), \
             patch.object(recalc.shutil, "which", return_value="/fixture/libreoffice"), \
             patch.object(recalc.subprocess, "run", side_effect=self.fake_conversion):
            report = recalc.recalculate(source, 5, output)
        self.assertEqual(report["status"], "success")
        self.assertEqual(report["quality_status"], "pending")
        self.assertTrue(report["source_unchanged"])
        self.assertEqual(output.read_bytes(), b"recalculated workbook")
        command = report["commands"][0]["command"]
        self.assertNotEqual(Path(command[-1]).parent, Path(command[command.index("--outdir") + 1]))
        self.assertTrue(any(x.startswith("-env:UserInstallation=file:") for x in command))

    def test_recalc_cli_report_collision_does_not_replace_source(self):
        source = self.source("working.xlsx")
        result = self.cli(RECALC_SCRIPT, source, "--report", source)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("Report path must differ", result.stderr)
        self.assertEqual(source.read_bytes(), b"unchanged fixture")

    @unittest.skipUnless(importlib.util.find_spec("xlsxwriter") and importlib.util.find_spec("openpyxl"),
                         "xlsxwriter and openpyxl are required for real formula fixtures")
    def test_real_xlsx_formula_errors_and_cache_inspection_are_read_only(self):
        import xlsxwriter
        source = self.root / "formula error.xlsx"
        with xlsxwriter.Workbook(source) as workbook:
            sheet = workbook.add_worksheet()
            sheet.write_formula("A1", "=1/0", None, "#DIV/0!")
        before = source.read_bytes()
        inspected = recalc.inspect(source)
        self.assertEqual(inspected["formula_count"], 1)
        self.assertEqual(inspected["errors"], [{"cell": "Sheet1!A1", "error": "#DIV/0!"}])
        report = verify.verify(source)
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any("#DIV/0!" in e for e in report["errors"]))
        self.assertEqual(source.read_bytes(), before)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is required for uncached formula fixture")
    def test_real_xlsx_missing_formula_caches_are_not_full_quality_pass(self):
        from openpyxl import Workbook
        source = self.root / "uncached.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "=1+1"
        workbook.save(source)
        workbook.close()
        before = source.read_bytes()
        inspected = recalc.inspect(source)
        self.assertTrue(inspected["missing_caches"])
        report = verify.verify(source)
        self.assertEqual(report["status"], "structural_pass")
        self.assertEqual(report["quality_status"], "pending")
        self.assertTrue(report["warnings"])
        self.assertEqual(source.read_bytes(), before)

    @unittest.skipUnless(importlib.util.find_spec("xlsxwriter") and importlib.util.find_spec("openpyxl")
                         and (shutil.which("libreoffice") or shutil.which("soffice")),
                         "LibreOffice, xlsxwriter and openpyxl are required for real manual-cache recalculation")
    def test_real_recalc_refreshes_stale_manual_mode_cache_without_mutating_source(self):
        import xlsxwriter
        from openpyxl import load_workbook
        source = self.root / "manual stale cache.xlsx"
        output = self.root / "recalculated.xlsx"
        with xlsxwriter.Workbook(source) as workbook:
            workbook.set_calc_mode("manual")
            workbook.add_worksheet().write_formula("A1", "=1+1", None, 999)
        before = source.read_bytes()
        original_values = load_workbook(source, data_only=True)
        self.assertEqual(original_values.active["A1"].value, 999)
        original_values.close()
        result = self.cli(RECALC_SCRIPT, source, "60", "--output", output, timeout=75)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        report = json.loads(result.stdout)
        self.assertEqual(report["status"], "success")
        self.assertTrue(report["source_unchanged"])
        self.assertEqual(source.read_bytes(), before)
        values = load_workbook(output, data_only=True)
        formulas = load_workbook(output, data_only=False)
        self.assertEqual(values.active["A1"].value, 2)
        self.assertEqual(formulas.active["A1"].value, "=1+1")
        values.close()
        formulas.close()

    @unittest.skipUnless(importlib.util.find_spec("xlsxwriter") and importlib.util.find_spec("openpyxl"),
                         "xlsxwriter and openpyxl are required for sparse worksheet fixture")
    def test_sparse_extreme_dimensions_fail_with_limit_instead_of_unbounded_iteration(self):
        import xlsxwriter
        source = self.root / "sparse dimensions.xlsx"
        with xlsxwriter.Workbook(source) as workbook:
            sheet = workbook.add_worksheet()
            sheet.write("A1", 1)
            sheet.write("XFD1048576", 2)
        before = source.read_bytes()
        verification = self.root / "verification.json"
        result = self.cli(SCRIPTS / "verify.py", source, "--output", verification, timeout=15)
        self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
        checked = json.loads(verification.read_text())
        self.assertEqual(checked["status"], "failed")
        self.assertRegex(" ".join(checked["errors"]), "(?i)(limit|budget|too large|exceeds)")
        output = self.root / "recalculated.xlsx"
        result = self.cli(RECALC_SCRIPT, source, "--output", output, timeout=15)
        self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
        report = json.loads(result.stdout)
        self.assertEqual(report["status"], "failed")
        self.assertRegex(report["error"], "(?i)(limit|budget|too large|exceeds)")
        self.assertEqual(report["commands"], [])
        self.assertFalse(output.exists())
        self.assertEqual(source.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
