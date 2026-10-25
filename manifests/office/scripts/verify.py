#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx", "openpyxl", "python-pptx", "pypdf"]
# ///
"""Reopen documents and inventory structure. This is NOT a visual/semantic QA pass."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import posixpath
import tempfile
from pathlib import Path
from urllib.parse import unquote
import xml.etree.ElementTree as ET
import zipfile


ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def sha256(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest() if hasattr(hashlib, "file_digest") else _hash(stream)


def _hash(stream) -> str:
    digest = hashlib.sha256()
    for chunk in iter(lambda: stream.read(1024 * 1024), b""):
        digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=path.parent, prefix=".office-", suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(data, stream, ensure_ascii=False, indent=2)
            stream.write("\n")
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def package_inventory(path: Path, required: str) -> dict:
    with zipfile.ZipFile(path) as package:
        entries = package.infolist()
        if len(entries) > 20000 or sum(x.file_size for x in entries) > 500 * 1024 * 1024:
            raise ValueError("Package exceeds inspection limits")
        names = package.namelist()
        if len(names) != len(set(names)):
            raise ValueError("Duplicate ZIP members")
        for info in entries:
            if info.filename.startswith("/") or ".." in Path(info.filename).parts or "\\" in info.filename:
                raise ValueError("Unsafe package member")
            if (info.external_attr >> 16) & 0o170000 == 0o120000:
                raise ValueError("Symlink package member")
        for member in ("[Content_Types].xml", "_rels/.rels", required):
            if member not in names:
                raise ValueError(f"Missing package part: {member}")
        bad = package.testzip()
        if bad:
            raise ValueError(f"ZIP checksum failure: {bad}")
        external = []
        for info in entries:
            if not info.filename.endswith((".xml", ".rels")):
                continue
            if info.file_size > 20 * 1024 * 1024:
                raise ValueError(f"XML inspection limit: {info.filename}")
            root = ET.fromstring(package.read(info))
            if info.filename.endswith(".rels"):
                base = posixpath.dirname(posixpath.dirname(info.filename))
                for rel in root:
                    target = rel.get("Target", "")
                    if rel.get("TargetMode") == "External":
                        external.append(target)
                        continue
                    target = unquote(target.split("#", 1)[0])
                    if not target:
                        continue
                    resolved = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(base, target))
                    if resolved not in names:
                        raise ValueError(f"Broken relationship: {info.filename} -> {target}")
        return {"parts": len(entries), "external_relationships": external}


def verify(path: Path, baseline: Path | None = None) -> dict:
    path = path.resolve()
    report = {"file": str(path), "format": path.suffix.lower().lstrip("."),
              "status": "failed", "quality_status": "pending", "errors": [], "warnings": [],
              "checks": {}, "coverage": {},
              "pending": ["Content/requirements review", "Full visual review", "Independent office-qa review"]}
    before = None
    try:
        before = sha256(path)
        report["sha256"] = before
        if baseline:
            report["baseline"] = {"file": str(baseline.resolve()), "sha256": sha256(baseline)}
        fmt = report["format"]
        word_formats = ("docx", "docm", "dotx", "dotm")
        if fmt in (*word_formats, "xlsx", "pptx"):
            required = "word/document.xml" if fmt in word_formats else {"xlsx": "xl/workbook.xml", "pptx": "ppt/presentation.xml"}[fmt]
            report["checks"]["package"] = package_inventory(path, required)
        if fmt in word_formats:
            # Count the actual package tree, including content controls, not
            # only top-level paragraphs exposed by a particular authoring API.
            with zipfile.ZipFile(path) as package:
                document = ET.fromstring(package.read("word/document.xml"))
                ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                report["coverage"] = {"paragraphs": len(list(document.iter(ns + "p"))),
                                      "tables": len(list(document.iter(ns + "tbl"))),
                                      "sections": len(list(document.iter(ns + "sectPr"))),
                                      "content_controls": len(list(document.iter(ns + "sdt")))}
                report["checks"]["text_characters"] = sum(len(n.text or "") for n in document.iter(ns + "t"))
                report["checks"]["macros_present"] = any(n.endswith("vbaProject.bin") for n in package.namelist())
                report["checks"]["macro_execution"] = "not_run"
        elif fmt == "pptx":
            from pptx import Presentation
            deck = Presentation(path)
            report["coverage"]["slides"] = len(deck.slides)
            report["checks"]["slides"] = [
                {"number": n, "shapes": len(slide.shapes),
                 "editable_text_shapes": sum(bool(shape.has_text_frame) for shape in slide.shapes),
                 "tables": sum(bool(shape.has_table) for shape in slide.shapes),
                 "charts": sum(bool(shape.has_chart) for shape in slide.shapes)}
                for n, slide in enumerate(deck.slides, 1)]
            if not deck.slides:
                report["errors"].append("Presentation has no slides")
        elif fmt == "xlsx":
            from openpyxl import load_workbook
            formulas = load_workbook(path, data_only=False)
            values = load_workbook(path, data_only=True)
            sheets = []
            for sheet in formulas:
                if sheet.max_row * sheet.max_column > 1_000_000:
                    raise ValueError(f"Worksheet inspection limit exceeded: {sheet.title}; review sparse/overformatted dimensions with a targeted workflow")
                count, missing = 0, []
                for row in sheet:
                    for cell in row:
                        cached = values[sheet.title][cell.coordinate]
                        if cell.data_type == "f":
                            count += 1
                            if cached.value is None:
                                missing.append(cell.coordinate)
                        if cached.data_type == "e":
                            report["errors"].append(f"Formula/cell error {sheet.title}!{cell.coordinate}: {cached.value}")
                sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                               "formulas": count, "empty_formula_caches": missing,
                               "charts": len(sheet._charts), "print_area": str(sheet.print_area)})
            report["coverage"]["worksheets"] = sheets
            if any(s["formulas"] for s in sheets):
                report["pending"].append("Recalculation evidence and independently checked expected values")
            if any(s["empty_formula_caches"] for s in sheets):
                report["warnings"].append("Empty formula caches: distinguish unrecalculated formulas from intentional empty strings")
            formulas.close()
            values.close()
        elif fmt == "pdf":
            from pypdf import PdfReader
            pdf = PdfReader(path, strict=True)
            if pdf.is_encrypted:
                raise ValueError("Encrypted PDF needs a separately approved password workflow")
            report["coverage"]["pages"] = len(pdf.pages)
            report["checks"]["text_characters"] = [len(p.extract_text() or "") for p in pdf.pages]
            if not pdf.pages:
                report["errors"].append("PDF has no pages")
            if any(x == 0 for x in report["checks"]["text_characters"]):
                report["warnings"].append("Pages without extractable text may need OCR; verify against source images")
        else:
            raise ValueError("Expected .docx/.docm/.dotx/.dotm, .xlsx, .pptx or .pdf")
        report["checks"]["reopened"] = True
    except Exception as exc:
        report["errors"].append(f"{type(exc).__name__}: {exc}")
    finally:
        if before is not None:
            try:
                report["checks"]["input_unchanged"] = sha256(path) == before
                if not report["checks"]["input_unchanged"]:
                    report["errors"].append("Input hash changed during inspection")
            except OSError as exc:
                report["errors"].append(str(exc))
    if not report["errors"]:
        report["status"] = "structural_pass"
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("--baseline", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.resolve() in {args.file.resolve(), args.baseline.resolve() if args.baseline else None}:
        parser.error("Report path must differ from input and baseline")
    report = verify(args.file, args.baseline)
    write_json(args.output, report)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "structural_pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
