#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx", "openpyxl", "xlsxwriter", "python-pptx", "pypdf", "Pillow"]
# ///
"""Generate real native Office fixtures and run installed helpers, never final QA."""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys

from verify import sha256, write_json


def fixtures(out: Path) -> dict[str, Path]:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches as WordInches
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt as SlidePt
    import xlsxwriter

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Noto Sans CJK SC"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    document.styles["Heading 1"].paragraph_format.keep_with_next = True
    title = document.add_heading("办公质量验证报告", 0)
    # An explicitly scoped native constraint: only the synthetic template title
    # is locked. Body text and tables remain editable; no global protection.
    control, props, content = OxmlElement("w:sdt"), OxmlElement("w:sdtPr"), OxmlElement("w:sdtContent")
    for name, value in (("tag", "office.template.title"), ("alias", "Synthetic template title"), ("lock", "sdtContentLocked")):
        prop = OxmlElement("w:" + name)
        prop.set(qn("w:val"), value)
        props.append(prop)
    title._p.addprevious(control)
    control.append(props)
    control.append(content)
    content.append(title._p)
    document.add_paragraph("此文件用于验证中文排版、原生段落、表格、页眉与页脚。所有数值均为合成测试数据。")
    document.add_heading("一、数据概览", 1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "项目", "销售额（元）"
    for label, number in (("一月", "100"), ("二月", "200"), ("三月", "300"), ("合计", "600")):
        row = table.add_row()
        row.cells[0].text, row.cells[1].text = label, number
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    # Reproducible companion visual and its editable SVG/data sources. The same
    # values drive the table, prose, spreadsheet and slide deck.
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (960, 280), "#F4F7F7")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("DejaVuSans.ttf", 24)
    for n, (month, amount) in enumerate((("Jan", 100), ("Feb", 200), ("Mar", 300))):
        y = 25 + 78 * n
        draw.text((24, y + 10), month, font=font, fill="#153B50")
        draw.rectangle((110, y, 110 + amount * 2, y + 52), fill="#217A8C")
        draw.text((130 + amount * 2, y + 10), str(amount), font=font, fill="#153B50")
    figure = out / "sales-figure.png"
    image.save(figure)
    svg = '<svg xmlns="http://www.w3.org/2000/svg" width="960" height="280" viewBox="0 0 960 280"><rect width="960" height="280" fill="#F4F7F7"/>'
    for n, (month, amount) in enumerate((("Jan", 100), ("Feb", 200), ("Mar", 300))):
        y = 25 + 78 * n
        svg += f'<text x="24" y="{y+34}" font-family="sans-serif" font-size="24">{month}</text><rect x="110" y="{y}" width="{amount*2}" height="52" fill="#217A8C"/><text x="{130+amount*2}" y="{y+34}" font-family="sans-serif" font-size="24">{amount}</text>'
    (out / "sales-figure.svg").write_text(svg + "</svg>", encoding="utf-8")
    write_json(out / "sales-data.json", {"source": "Synthetic smoke fixture, not observed business data", "unit": "CNY", "rows": [["Jan", 100], ["Feb", 200], ["Mar", 300]]})
    picture = document.add_picture(str(figure), width=WordInches(5.7))
    picture._inline.docPr.set("descr", "合成月销售额：一月100元、二月200元、三月300元；逐月上升。")
    caption = document.add_paragraph(style="Caption")
    caption.add_run("图 ")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "17")
    start.set(qn("w:name"), "SalesFigureNumber")
    caption._p.append(start)
    seq = OxmlElement("w:fldSimple")
    seq.set(qn("w:instr"), " SEQ Figure \\* ARABIC ")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    seq.append(run)
    caption._p.append(seq)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "17")
    caption._p.append(end)
    caption.add_run("  月销售趋势（合成测试数据，单位：元）")
    reference = document.add_paragraph("如图 ")
    ref = OxmlElement("w:fldSimple")
    ref.set(qn("w:instr"), " REF SalesFigureNumber \\h ")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    ref.append(run)
    reference._p.append(ref)
    reference.add_run(" 所示，销售额逐月上升；三个月合计600元，月均200元。")
    document.sections[0].header.paragraphs[0].text = "OFFICE · 合成数据验收"
    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "第 "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.add_page_break()
    document.add_heading("二、交付检查", 1)
    for text in ("输入副本与源文件哈希留痕。", "核对标题层级、表格边框和分页。", "结构校验、渲染与独立审查分别记录，不相互替代。"):
        document.add_paragraph(text, style="List Bullet")
    document.add_paragraph("结论：本报告是测试样本，最终质量状态由独立审阅给出。")
    docx = out / "中文报告.docx"
    document.save(docx)

    original = out / "source-model.xlsx"
    workbook = xlsxwriter.Workbook(original, {"strings_to_formulas": False, "strings_to_urls": False})
    header = workbook.add_format({"bold": True, "bg_color": "#153B50", "font_color": "white", "font_name": "Noto Sans CJK SC"})
    money = workbook.add_format({"num_format": "#,##0", "font_name": "Noto Sans CJK SC"})
    data = workbook.add_worksheet("Sales Data")
    summary = workbook.add_worksheet("Summary")
    for sheet in (data, summary):
        sheet.set_column("A:A", 22)
        sheet.set_column("B:C", 18)
        sheet.set_paper(9)
        sheet.fit_to_pages(1, 1)
        sheet.set_landscape()
    data.write_row("A1", ["月份", "销售额（元）"], header)
    for n, (month, amount) in enumerate((("一月", 100), ("二月", 200), ("三月", 300)), 1):
        data.write(n, 0, month)
        data.write(n, 1, amount, money)
    data.write("A5", "合计", header)
    data.write_formula("B5", "=SUM(B2:B4)", money, 600)
    data.print_area("A1:C7")
    summary.write_row("A1", ["指标", "数值"], header)
    summary.write("A2", "销售总额")
    summary.write_formula("B2", "='Sales Data'!B5", money, 600)
    summary.write("A3", "月均销售额")
    summary.write_formula("B3", "=B2/3", money, 200)
    chart = workbook.add_chart({"type": "column"})
    chart.add_series({"name": "月销售额", "categories": "='Sales Data'!$A$2:$A$4", "values": "='Sales Data'!$B$2:$B$4", "fill": {"color": "#217A8C"}})
    chart.set_title({"name": "销售趋势（合成数据）"})
    chart.set_legend({"none": True})
    summary.insert_chart("A6", chart, {"x_scale": 0.85, "y_scale": 0.85})
    summary.print_area("A1:H22")
    workbook.close()

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for title, subtitle in (("办公质量验证", "原生可编辑 · 中文排版 · 合成数据"),
                             ("销售结果", "三个月合计 600 元，平均每月 200 元"),
                             ("验收结论", "结构、数值、视觉和独立 QA 分别留痕")):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F4F7F7")
        for text, top, size in ((title, 0.6, 32), (subtitle, 1.55, 20)):
            box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(0.8))
            paragraph = box.text_frame.paragraphs[0]
            paragraph.text = text
            paragraph.font.name, paragraph.font.size = "Noto Sans CJK SC", SlidePt(size)
            paragraph.font.color.rgb = RGBColor.from_string("153B50")
    chart_data = CategoryChartData()
    chart_data.categories = ["一月", "二月", "三月"]
    chart_data.add_series("销售额", [100, 200, 300])
    deck.slides[1].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(2.6), Inches(6.8), Inches(3.8), chart_data)
    table = deck.slides[1].shapes.add_table(3, 2, Inches(8.2), Inches(2.9), Inches(4.1), Inches(2.4)).table
    for row, values in enumerate((("指标", "数值"), ("总额", "600 元"), ("月均", "200 元"))):
        for col, text in enumerate(values):
            table.cell(row, col).text = text
    for n, text in enumerate(("数据一致", "文件可编辑", "逐页审阅")):
        shape = deck.slides[2].shapes.add_shape(1, Inches(0.8 + n * 4.1), Inches(3.1), Inches(3.6), Inches(1.7))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string("217A8C")
        shape.text = text
        shape.text_frame.paragraphs[0].font.size = SlidePt(24)
    pptx = out / "可编辑演示.pptx"
    deck.save(pptx)
    return {"docx": docx, "xlsx_source": original, "pptx": pptx}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()
    out = args.out_dir.resolve()
    if out.exists() and any(out.iterdir()):
        parser.error("Use a fresh output directory")
    out.mkdir(parents=True, exist_ok=True)
    report = {"status": "failed", "quality_status": "pending", "commands": [], "artifacts": {},
              "note": "Synthetic smoke fixtures; independent visual QA is still required."}
    scripts = Path(__file__).resolve().parent

    def execute(command: list[str]) -> None:
        result = subprocess.run(command, capture_output=True, text=True, timeout=240,
                                env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
        report["commands"].append({"command": command, "exit_status": result.returncode,
                                   "stdout": result.stdout, "stderr": result.stderr})
        if result.returncode:
            raise ValueError(f"Command failed: {command}; see smoke.json")

    def check_render(file: Path, name: str) -> None:
        execute([sys.executable, str(scripts / "verify.py"), str(file), "--output", str(out / f"{name}-structure.json")])
        execute([sys.executable, str(scripts / "render.py"), str(file), "--out-dir", str(out / f"{name}-preview")])
        report["artifacts"][name] = {"file": str(file), "sha256": sha256(file),
                                    "verification": str(out / f"{name}-structure.json"), "render": str(out / f"{name}-preview/render.json")}

    try:
        files = fixtures(out)
        source_hash = sha256(files["xlsx_source"])
        working = out / "销售模型.xlsx"
        shutil.copy2(files["xlsx_source"], working)
        execute([sys.executable, str(scripts.parent / "xlsx/scripts/recalc.py"), str(working), "--report", str(out / "recalc.json")])
        from openpyxl import load_workbook
        values = load_workbook(working, data_only=True)
        assert values["Sales Data"]["B5"].value == 600, "Incorrect total"
        assert values["Summary"]["B2"].value == 600, "Incorrect cross-sheet reference"
        assert values["Summary"]["B3"].value == 200, "Incorrect mean"
        values.close()
        assert sha256(files["xlsx_source"]) == source_hash, "Source workbook changed"
        report["expected_values"] = {"Sales Data!B5": 600, "Summary!B2": 600, "Summary!B3": 200}
        report["source_unchanged"] = True
        check_render(files["docx"], "docx")
        expected = {"styles": {"Heading1": {"keep_with_next": True}},
                    "content_controls": {"office.template.title": {"lock": "sdtContentLocked"}},
                    "fields": ["SEQ", "REF", "PAGE"],
                    "figures": {"min_count": 1, "require_alt_text": True},
                    "tables": {"min_count": 1, "repeat_header": True, "prevent_row_split": True},
                    "document_protection": {"enforced": False}}
        write_json(out / "native-constraints.json", expected)
        execute([sys.executable, str(scripts.parent / "docx/scripts/feature_audit.py"), str(files["docx"]),
                 "--expect", str(out / "native-constraints.json"), "--output", str(out / "docx-native-features.json")])
        report["docx_native_features"] = str(out / "docx-native-features.json")
        report["docx_visual_sources"] = [str(out / name) for name in ("sales-data.json", "sales-figure.svg", "sales-figure.png")]
        check_render(working, "xlsx")
        check_render(files["pptx"], "pptx")
        from pypdf import PdfReader, PdfWriter
        source_pdf = out / "docx-preview" / (files["docx"].stem + ".pdf")
        source_reader = PdfReader(source_pdf)
        assert len(source_reader.pages) == 2, "DOCX should render as two pages"
        merged = out / "合并报告.pdf"
        writer = PdfWriter()
        writer.append(source_pdf)
        writer.append(source_pdf)
        with merged.open("wb") as stream:
            writer.write(stream)
        assert len(PdfReader(merged).pages) == 4
        reader = PdfReader(merged)
        for name, indices in (("pdf-first", [0]), ("pdf-rest", [1, 2, 3])):
            writer = PdfWriter()
            for index in indices:
                writer.add_page(reader.pages[index])
            split = out / f"{name}.pdf"
            with split.open("wb") as stream:
                writer.write(stream)
            assert len(PdfReader(split).pages) == len(indices)
            check_render(split, name)
        check_render(merged, "pdf-merged")
        report.update(status="smoke_pass", pdf_merge_pages=4, pdf_split_pages=[1, 3])
    except Exception as exc:
        report["error"] = f"{type(exc).__name__}: {exc}"
    write_json(out / "smoke.json", report)
    print(json.dumps({key: value for key, value in report.items() if key != "commands"}, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "smoke_pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
