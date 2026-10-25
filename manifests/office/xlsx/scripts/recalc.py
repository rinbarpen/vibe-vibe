#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Recalculate a working XLSX copy; validate before atomic replacement."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def inspect(path: Path) -> dict:
    from openpyxl import load_workbook
    with zipfile.ZipFile(path) as package:
        members = package.infolist()
        if len(members) > 20000 or sum(x.file_size for x in members) > 500 * 1024 * 1024:
            raise ValueError("Workbook package inspection limit exceeded")
        if any(x.file_size > 20 * 1024 * 1024 for x in members if x.filename.endswith(".xml")):
            raise ValueError("Workbook XML inspection limit exceeded; use a targeted large-workbook workflow")
        if any(name.startswith("xl/externalLinks/") or name.endswith("vbaProject.bin")
               for name in package.namelist()):
            raise ValueError("External links or VBA need an application-specific preservation workflow")
        missing_caches = []
        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        for name in package.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                root = ET.fromstring(package.read(name))
                for cell in root.findall(".//x:c", ns):
                    if cell.find("x:f", ns) is not None:
                        cached = cell.find("x:v", ns)
                        if cached is None or (cached.text is None and cell.get("t") != "str"):
                            missing_caches.append(f"{name}!{cell.get('r')}")
    workbook = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    formulas, errors, charts = {}, [], 0
    try:
        for sheet in workbook:
            if sheet.max_row * sheet.max_column > 1_000_000:
                raise ValueError(f"Worksheet inspection limit exceeded: {sheet.title}; use a targeted large/sparse-workbook workflow")
            charts += len(sheet._charts)
            for row in sheet:
                for cell in row:
                    key = f"{sheet.title}!{cell.coordinate}"
                    if cell.data_type == "f":
                        if not isinstance(cell.value, str):
                            raise ValueError("Array/data-table formula needs application-specific verification")
                        formulas[key] = cell.value
                    cached = values[sheet.title][cell.coordinate]
                    if cached.data_type == "e":
                        errors.append({"cell": key, "error": cached.value})
        return {"formulas": formulas, "formula_count": len(formulas), "errors": errors,
                "missing_caches": missing_caches, "sheet_names": workbook.sheetnames,
                "charts": charts}
    finally:
        workbook.close()
        values.close()


def prepare_recalculation(source: Path, destination: Path) -> None:
    """On the TEMPORARY copy only, force auto mode and remove stale caches.

    A manual-calculation workbook can round-trip through LibreOffice with its
    old cache intact. Removing caches makes a no-op conversion fail verification
    rather than incorrectly proving recalculation.
    """
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", ns)
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as outgoing:
        for info in incoming.infolist():
            data = incoming.read(info)
            if info.filename == "xl/workbook.xml":
                root = ET.fromstring(data)
                calc = root.find(f"{{{ns}}}calcPr")
                if calc is None:
                    calc = ET.SubElement(root, f"{{{ns}}}calcPr")
                calc.set("calcMode", "auto")
                calc.set("fullCalcOnLoad", "1")
                calc.set("forceFullCalc", "1")
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                root = ET.fromstring(data)
                for cell in root.iter(f"{{{ns}}}c"):
                    if cell.find(f"{{{ns}}}f") is not None:
                        for cached in list(cell.findall(f"{{{ns}}}v")):
                            cell.remove(cached)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            outgoing.writestr(info, data)


def recalculate(source: Path, timeout: int, output: Path | None = None) -> dict:
    source = source.resolve()
    output = output.resolve() if output else source
    report = {"status": "failed", "source": str(source), "output": str(output),
              "quality_status": "pending", "commands": [],
              "note": "Evaluated formulas still require independent expected-value checks and visual QA."}
    try:
        before = digest(source)
        report["source_sha256_before"] = before
        if timeout <= 0 or source.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
            raise ValueError("Use XLSX paths and a positive timeout")
        original = inspect(source)
        report["before"] = original
        office = shutil.which("libreoffice") or shutil.which("soffice")
        if not office:
            raise ValueError("Missing LibreOffice")
        with tempfile.TemporaryDirectory(prefix="office-recalc-") as directory:
            work = Path(directory)
            incoming, outgoing = work / "input", work / "output"
            incoming.mkdir()
            outgoing.mkdir()
            copied = incoming / source.name
            prepare_recalculation(source, copied)
            report["forced_recalculation"] = "auto mode; removed formula caches on temporary input only"
            command = [office, f"-env:UserInstallation={(work / 'profile').as_uri()}", "--headless",
                       "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(outgoing), str(copied)]
            try:
                result = subprocess.run(command, capture_output=True, text=True, timeout=timeout)
            except subprocess.TimeoutExpired as exc:
                report["commands"].append({"command": command, "exit_status": None, "timed_out": True,
                                           "stdout": str(exc.stdout or ""), "stderr": str(exc.stderr or "")})
                raise ValueError(f"Recalculation timed out after {timeout}s") from exc
            report["commands"].append({"command": command, "exit_status": result.returncode,
                                       "stdout": result.stdout, "stderr": result.stderr})
            if result.returncode:
                raise ValueError(f"LibreOffice exited {result.returncode}")
            candidate = outgoing / source.name
            if not candidate.is_file() or candidate.stat().st_size == 0:
                raise ValueError("LibreOffice returned without creating a new workbook")
            actual = inspect(candidate)
            report["after"] = actual
            if actual["formulas"] != original["formulas"]:
                raise ValueError("Recalculation changed formula text or locations; original working copy retained")
            if actual["sheet_names"] != original["sheet_names"] or actual["charts"] != original["charts"]:
                raise ValueError("Recalculation changed worksheet/chart inventory")
            if actual["errors"] or actual["missing_caches"]:
                raise ValueError("Recalculated workbook contains errors or absent formula caches")
            if digest(source) != before:
                raise ValueError("Source changed concurrently; replacement cancelled")
            output.parent.mkdir(parents=True, exist_ok=True)
            fd, temporary = tempfile.mkstemp(dir=output.parent, prefix=".recalc-", suffix=".xlsx")
            try:
                with os.fdopen(fd, "wb") as stream, candidate.open("rb") as data:
                    shutil.copyfileobj(data, stream)
                    stream.flush()
                    os.fsync(stream.fileno())
                os.chmod(temporary, source.stat().st_mode & 0o777)
                os.replace(temporary, output)
            finally:
                if os.path.exists(temporary):
                    os.unlink(temporary)
        report.update(status="success", output_sha256=digest(output), total_errors=0,
                      total_formulas=actual["formula_count"])
    except Exception as exc:
        report["error"] = f"{type(exc).__name__}: {exc}"
    if source.is_file():
        report["source_sha256_after"] = digest(source)
        report["source_unchanged"] = report.get("source_sha256_before") == report["source_sha256_after"]
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("timeout", type=int, nargs="?", default=120)
    parser.add_argument("--output", type=Path, help="Optional different XLSX destination")
    parser.add_argument("--report", type=Path, help="Optional JSON record")
    args = parser.parse_args()
    if args.report and args.report.resolve() in {args.file.resolve(), args.output.resolve() if args.output else None}:
        parser.error("Report path must differ from workbook paths")
    report = recalculate(args.file, args.timeout, args.output)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        fd, temporary = tempfile.mkstemp(dir=args.report.parent, prefix=".recalc-report-")
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as stream:
                json.dump(report, stream, ensure_ascii=False, indent=2)
                stream.write("\n")
            os.replace(temporary, args.report)
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "success" else 1


if __name__ == "__main__":
    raise SystemExit(main())
